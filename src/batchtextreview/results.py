"""Turning a RunResult into tables, exports and history records.

Shared by the UI and the CLI so both produce identical files.
"""

from __future__ import annotations

import datetime as _dt
import io
import json
import math
import re
from collections.abc import Sequence
from decimal import Decimal
from typing import TYPE_CHECKING, Any, Literal

import numpy as np
import pandas as pd

from batchtextreview import __version__
from batchtextreview import labels as _labels

if TYPE_CHECKING:
    from batchtextreview.engine import RunResult
    from batchtextreview.evaluation import EvalReport

OUTPUT_COL = "OUTPUT"
LABEL_COL = "LABEL"
STATUS_COL = "STATUS"
ERROR_COL = "ERROR"
NOTE_COL = "NOTE"

NOTE_INPUT_TRUNCATED = "input truncated"
NOTE_OUTPUT_LIMIT = "output may be cut off"

SCORABLE_STATUSES = frozenset({"generated", "cached"})

JSON_VALID_SUFFIX = "valid_json"
MAX_JSON_COLUMNS = 50
"""At most this many ``OUTPUT.<key>`` columns (the most frequent keys win),
so a model that answers with free-form keys can't add thousands of columns."""

CSV_MIME = "text/csv"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EXCEL_MAX_CELL_CHARS = 32767
"""Excel refuses (and "repairs") cells longer than this."""
EXCEL_MAX_COL_WIDTH = 80
_EXCEL_MIN_COL_WIDTH = 8
_EXCEL_MAX_EXACT_INT = 10**15 - 1
_WIDTH_SAMPLE_ROWS = 1000


def free_column_name(frame: pd.DataFrame, wanted: str) -> str:
    """``wanted`` if unused in ``frame``, else ``wanted_2``, ``wanted_3``..."""
    return _free_name({str(c) for c in frame.columns}, wanted)


def scorable_mask(result: RunResult) -> list[bool]:
    """True for rows with a real model output (generated or cached)."""
    return [status in SCORABLE_STATUSES for status in result.statuses]


def build_results_frame(
    frame: pd.DataFrame,
    result: RunResult,
    labels: Sequence[str] | None = None,
    expand_json: bool = False,
) -> pd.DataFrame:
    """The working rows plus result columns, original order and columns kept.

    ``frame`` has exactly ``len(result.outputs)`` rows (ValueError otherwise);
    its index is reset in the output. Adds, using :func:`free_column_name`
    so existing columns are never overwritten:

    * OUTPUT - the cleaned output ("" for blank/error/pending rows);
    * LABEL - when ``labels`` is given: ``labels.resolve_labels`` (canonical
      label, INVALID, or "" for non-scorable rows);
    * STATUS - the row status;
    * ERROR - only when at least one row errored;
    * NOTE - only when at least one row's text was cut to
      ``max_input_chars`` ("input truncated") or its output reached
      ``max_new_tokens`` ("output may be cut off"); both are joined with
      "; ". Per-row flags of the wrong length are ignored.
    * With ``expand_json``: one extra column per top-level key of JSON
      object outputs (prefixed "OUTPUT.", e.g. "OUTPUT.people"); list values
      joined with "; ", nested objects as compact JSON, unparsable outputs
      leave the cells empty. Adds "OUTPUT.valid_json" (bool) too.

    Details: a scorable row whose output is empty gets LABEL = INVALID (the
    evaluation scores it that way too). The JSON columns use the actual
    OUTPUT column name as prefix (``OUTPUT_2.people`` after a collision);
    ``valid_json`` is True/False for scorable rows and empty (None) for
    blank / error / pending rows, which have no output to judge. At most
    :data:`MAX_JSON_COLUMNS` key columns are added.
    """
    n = len(result.outputs)
    if len(frame) != n:
        raise ValueError(f"The table has {len(frame):,} rows but the run produced {n:,} outputs.")
    table = frame.reset_index(drop=True)
    taken = {str(c) for c in table.columns}
    new_columns: dict[str, pd.Series] = {}

    def add(wanted: str, values: Sequence[Any]) -> str:
        name = _free_name(taken, wanted)
        taken.add(name)
        new_columns[name] = pd.Series(list(values), index=table.index, dtype=object)
        return name

    scorable = scorable_mask(result)
    outputs = ["" if _labels.is_missing(o) else str(o) for o in result.outputs]
    statuses = [str(s) for s in result.statuses]

    output_col = add(OUTPUT_COL, outputs)
    if labels:
        resolved = _labels.resolve_labels(
            [o if keep else "" for o, keep in zip(outputs, scorable, strict=True)], labels
        )
        add(
            LABEL_COL,
            [
                (label or _labels.INVALID) if keep else ""
                for label, keep in zip(resolved, scorable, strict=True)
            ],
        )
    add(STATUS_COL, statuses)
    if any(s == "error" for s in statuses):
        errors = list(result.errors) if len(result.errors) == n else [None] * n
        add(
            ERROR_COL,
            [
                (str(e) if e else "error") if s == "error" else ""
                for e, s in zip(errors, statuses, strict=True)
            ],
        )
    notes = _row_notes(result, n)
    if any(notes):
        add(NOTE_COL, notes)
    if expand_json:
        _add_json_columns(add, taken, output_col, outputs, scorable)

    # One concat instead of column-by-column inserts: no fragmentation
    # warnings, and duplicate column names in ``frame`` are left alone.
    return pd.concat([table, pd.DataFrame(new_columns, index=table.index)], axis=1)


def to_csv_bytes(frame: pd.DataFrame) -> bytes:
    """CSV encoded as UTF-8 *with BOM* so Excel on Windows shows non-ASCII
    text correctly. No index.

    Text cells, header included, that a spreadsheet would run as a formula
    (leading ``=``, ``@``, tab or carriage return, or ``+`` / ``-`` not
    starting a plain number) are prefixed with a single quote. Numbers such
    as "-5" or "+3.5" are left as they are."""
    safe = frame.copy()
    for position in range(safe.shape[1]):
        column = safe.iloc[:, position]
        if _is_texty(column.dtype):
            safe.isetitem(position, column.astype(object).map(_neutralise_formula))
    if not isinstance(safe.columns, pd.MultiIndex):
        # Header cells too (a new Index keeps duplicate names).
        safe.columns = pd.Index([_neutralise_formula(c) for c in safe.columns], dtype=object)
    text = safe.to_csv(index=False, lineterminator="\r\n")
    return text.encode("utf-8-sig", errors="replace")


def to_excel_bytes(frame: pd.DataFrame, sheet_name: str = "results") -> bytes:
    """An .xlsx (openpyxl) with a frozen header row and sensible column
    widths (max 80 chars). No index.

    Text is always stored as text - a cell starting with "=" never becomes
    a formula, nor does "#N/A" (or another error code) become an error
    value. Characters XML can't hold are dropped and cells are cut at
    Excel's 32,767-character limit."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ERROR_CODES
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=_sheet_title(sheet_name))
    sheet.freeze_panes = "A2"
    headers = [_excel_text(str(c)) for c in frame.columns]
    for position, width in enumerate(_column_widths(frame, headers), start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="F1F3F5")
    header_row = []
    for name in headers:
        cell = WriteOnlyCell(sheet, value=name)
        cell.data_type = "s"
        cell.font = header_font
        cell.fill = header_fill
        header_row.append(cell)
    sheet.append(header_row)

    for values in frame.itertuples(index=False, name=None):
        row: list[Any] = []
        for value in values:
            converted = _excel_value(value)
            if isinstance(converted, str) and (
                converted.startswith("=") or converted in ERROR_CODES
            ):
                # openpyxl infers a formula from a leading "=" and an error
                # cell from text such as "#N/A"; force text.
                cell = WriteOnlyCell(sheet, value=converted)
                cell.data_type = "s"
                row.append(cell)
            else:
                row.append(converted)
        sheet.append(row)

    if headers:
        last = get_column_letter(len(headers))
        sheet.auto_filter.ref = f"A1:{last}{len(frame) + 1}"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


ExportFormat = Literal["csv", "xlsx"]


def export_bytes(frame: pd.DataFrame, fmt: ExportFormat) -> tuple[bytes, str]:
    """(payload, mime type) for the given format."""
    if fmt == "csv":
        return to_csv_bytes(frame), CSV_MIME
    if fmt == "xlsx":
        return to_excel_bytes(frame), XLSX_MIME
    raise ValueError(f"Unknown export format {fmt!r}; expected 'csv' or 'xlsx'.")


def history_record(
    report: EvalReport,
    result: RunResult,
    *,
    preset_key: str,
    var: str | None,
    dataset: str | None,
    dataset_fingerprint: str | None,
    text_column: str,
    gold_column: str,
    n_rows: int,
    source: str = "app",
    selection: str | None = None,
    template_edited: bool | None = None,
    preset_title: str | None = None,
) -> dict[str, Any]:
    """The dict passed to ``Store.log_evaluation``: run/model fields from
    ``result.config``, metrics from ``report`` (accuracy, macro_f1, kappa,
    n_scored, n_correct, n_invalid=n_pred_invalid, n_excluded,
    n_gold_unscorable), ``labels`` = report.labels in labels mode else None.
    Floats rounded to 4 dp.

    Provenance: ``selection`` (which rows ran, e.g. "First 100 of 2,000
    rows"), ``template_edited`` (the prompt differs from the preset's) and
    ``preset_title`` are passed through; ``rows_per_s`` (generated rows per
    second of generation), ``elapsed_s`` and ``app_version`` are filled in.
    ``rows_per_s`` and ``elapsed_s`` are None when nothing was generated: a
    run served from the cache says nothing about speed, and re-scoring it
    must not overwrite the figures its first run logged (``log_evaluation``
    keeps them when given None)."""
    config = result.config
    stats = result.stats
    rows_per_s = (
        stats.generated / stats.generation_s if stats.generated and stats.generation_s > 0 else None
    )
    elapsed_s = stats.elapsed_s if stats.generated else None
    return {
        "run_id": result.run_id,
        "model_id": config.model_id,
        "model_label": config.model_label,
        "precision": config.precision,
        "preset": preset_key,
        "var": (var or "").strip() or None,
        "labels": list(report.labels) if report.mode == "labels" else None,
        "dataset": dataset,
        "dataset_fingerprint": dataset_fingerprint,
        "text_column": text_column,
        "gold_column": gold_column,
        "n_rows": int(n_rows),
        "n_scored": int(report.n_scored),
        "n_correct": int(report.n_correct),
        "accuracy": _round4(report.accuracy),
        "macro_f1": _round4(report.macro_f1),
        "kappa": _round4(report.kappa),
        "n_invalid": int(report.n_pred_invalid),
        "n_excluded": int(report.n_excluded),
        "n_gold_unscorable": int(report.n_gold_unscorable),
        "source": source or "app",
        "selection": (selection or "").strip() or None,
        "rows_per_s": _round4(rows_per_s),
        "elapsed_s": _round4(elapsed_s),
        "app_version": __version__,
        "template_edited": None if template_edited is None else bool(template_edited),
        "preset_title": (preset_title or "").strip() or None,
    }


def summarise_stats(result: RunResult, *, token_setting: str = "Max output tokens") -> str:
    """One human sentence, e.g. "120 rows: 80 generated, 30 from cache,
    10 blank. 12 duplicate rows were computed once." Mentions errors /
    pending / cancellation / abort when present, and texts cut to
    ``max_input_chars`` / outputs that reached ``max_new_tokens`` (with
    advice to raise ``token_setting``: the app's label by default, the
    CLI passes its option). Uses singular/plural correctly and omits zero
    parts."""
    stats = result.stats
    parts = [
        f"{count:,} {what}"
        for count, what in (
            (stats.generated, "generated"),
            (stats.cached, "from cache"),
            (stats.blank, "blank"),
            (stats.errors, "failed"),
            (stats.pending, "not run"),
        )
        if count
    ]
    rows = _plural(stats.total, "row")
    sentences = [f"{rows}: {', '.join(parts)}." if parts else f"{rows}."]
    if stats.duplicates:
        verb = "was" if stats.duplicates == 1 else "were"
        sentences.append(f"{_plural(stats.duplicates, 'duplicate row')} {verb} computed once.")
    if stats.truncated_inputs:
        one = stats.truncated_inputs == 1
        sentences.append(
            f"{_plural(stats.truncated_inputs, 'text')} {'was' if one else 'were'} longer than "
            f"{result.config.max_input_chars:,} characters and {'was' if one else 'were'} cut "
            "before sending."
        )
    if stats.hit_token_limit:
        sentences.append(
            f"{_plural(stats.hit_token_limit, 'output')} reached the "
            f"{result.config.max_new_tokens:,}-token limit and may be cut off - raise "
            f"{token_setting}."
        )
    if result.cancelled:
        sentences.append("The run was cancelled.")
    if result.aborted_reason:
        reason = " ".join(result.aborted_reason.split())
        if len(reason) > 200:
            reason = reason[:199].rstrip() + "…"
        sentences.append(f"Stopped early: {reason.rstrip('.')}.")
    return " ".join(sentences)


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------


def _free_name(taken: set[str], wanted: str) -> str:
    if wanted not in taken:
        return wanted
    n = 2
    while f"{wanted}_{n}" in taken:
        n += 1
    return f"{wanted}_{n}"


def _plural(n: int, noun: str) -> str:
    return f"{n:,} {noun}" if n == 1 else f"{n:,} {noun}s"


def _row_notes(result: RunResult, n: int) -> list[str]:
    """Per-row NOTE text ("" for most rows)."""

    def flags(values: Sequence[Any]) -> list[bool]:
        return [bool(v) for v in values] if len(values) == n else [False] * n

    notes = []
    for cut_input, cut_output in zip(
        flags(result.input_truncated), flags(result.hit_token_limit), strict=True
    ):
        parts = [NOTE_INPUT_TRUNCATED] if cut_input else []
        if cut_output:
            parts.append(NOTE_OUTPUT_LIMIT)
        notes.append("; ".join(parts))
    return notes


def _round4(value: Any) -> float | None:
    if value is None:
        return None
    number = float(value)
    return None if math.isnan(number) else round(number, 4)


# ---- JSON expansion ---------------------------------------------------------


def _add_json_columns(
    add: Any,
    taken: set[str],
    prefix: str,
    outputs: Sequence[str],
    scorable: Sequence[bool],
) -> None:
    parsed = [
        _labels.parse_json_output(text) if keep and text else None
        for text, keep in zip(outputs, scorable, strict=True)
    ]
    # Reserve the flag's name first so a JSON key called "valid_json" can't
    # take it.
    valid_name = _free_name(taken, f"{prefix}.{JSON_VALID_SUFFIX}")
    taken.add(valid_name)

    frequency: dict[str, int] = {}
    for value in parsed:
        if isinstance(value, dict):
            for key in value:
                frequency[str(key)] = frequency.get(str(key), 0) + 1
    keys = list(frequency)  # first-seen order
    if len(keys) > MAX_JSON_COLUMNS:
        keep = set(sorted(keys, key=lambda k: -frequency[k])[:MAX_JSON_COLUMNS])
        keys = [k for k in keys if k in keep]

    for key in keys:
        cells = []
        for value in parsed:
            if isinstance(value, dict) and key in value:
                cells.append(_json_cell(value[key]))
            else:
                cells.append("")
        add(f"{prefix}.{key}", cells)

    taken.discard(valid_name)
    flags = [
        (value is not None) if keep else None for value, keep in zip(parsed, scorable, strict=True)
    ]
    add(valid_name, flags)


def _json_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        items = (_json_scalar(item) for item in value if item is not None)
        return "; ".join(item for item in items if item)
    return _json_scalar(value)


def _json_scalar(value: Any) -> str:
    if isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


# ---- CSV ------------------------------------------------------------------

_PLAIN_NUMBER = re.compile(r"[+-]?(?:\d[\d,]*(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?%?")


def _is_texty(dtype: Any) -> bool:
    return (
        pd.api.types.is_object_dtype(dtype)
        or isinstance(dtype, (pd.StringDtype, pd.CategoricalDtype))
        or pd.api.types.is_string_dtype(dtype)
    )


def _neutralise_formula(value: Any) -> Any:
    if not isinstance(value, str) or not value:
        return value
    first = value[0]
    if first in "=@\t\r":
        return "'" + value
    # A lone sign or a plain number ("-5", "+3.5") is harmless.
    if first in "+-" and len(value) > 1 and not _PLAIN_NUMBER.fullmatch(value):
        return "'" + value
    return value


# ---- Excel ----------------------------------------------------------------

_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


def _sheet_title(name: str) -> str:
    title = _SHEET_FORBIDDEN.sub(" ", str(name or "")).strip().strip("'")[:31].strip()
    return title or "results"


def _excel_text(text: str) -> str:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    try:
        text.encode("utf-8")
    except UnicodeEncodeError:  # lone surrogates from a lenient decode
        text = text.encode("utf-8", errors="replace").decode("utf-8")
    if len(text) > EXCEL_MAX_CELL_CHARS:
        text = text[: EXCEL_MAX_CELL_CHARS - 1] + "…"
    return text


def _excel_value(value: Any) -> Any:
    """A value openpyxl can write faithfully (None for missing)."""
    if value is None:
        return None
    if isinstance(value, str):
        return _excel_text(value)
    if isinstance(value, np.datetime64):
        value = pd.Timestamp(value)
    elif isinstance(value, np.timedelta64):
        value = pd.Timedelta(value)
    elif isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        # Excel keeps 15 significant digits: long ids would be silently
        # rounded (1234567890123456789 -> 1.23457E+18), so write them as text.
        return value if abs(value) <= _EXCEL_MAX_EXACT_INT else str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return None
        if math.isinf(value):
            return "inf" if value > 0 else "-inf"
        return value
    if isinstance(value, Decimal):
        return value if value.is_finite() else str(value)
    if _labels.is_missing(value):  # pd.NA, NaT
        return None
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, (_dt.datetime, _dt.time)):
        return value.replace(tzinfo=None) if value.tzinfo is not None else value
    if isinstance(value, _dt.date):
        return value
    if isinstance(value, (dict, list, tuple)):
        try:
            return _excel_text(json.dumps(value, ensure_ascii=False, default=str))
        except (TypeError, ValueError):
            pass
    return _excel_text(str(value))


def _column_widths(frame: pd.DataFrame, headers: Sequence[str]) -> list[float]:
    sample = frame.head(_WIDTH_SAMPLE_ROWS)
    widths: list[float] = []
    for position, header in enumerate(headers):
        longest = len(header)
        for value in sample.iloc[:, position]:
            if _labels.is_missing(value):
                continue
            first_line = str(value).split("\n", 1)[0]
            longest = max(longest, len(first_line))
            if longest >= EXCEL_MAX_COL_WIDTH:
                break
        widths.append(float(min(EXCEL_MAX_COL_WIDTH, max(_EXCEL_MIN_COL_WIDTH, longest + 2))))
    return widths
