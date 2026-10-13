"""Tests for batchtextreview.data: reading files and guessing column roles."""

from __future__ import annotations

import codecs
import io
import json
import logging
import os
import re
from pathlib import Path

import pandas as pd
import pytest

from batchtextreview import results
from batchtextreview.data import (
    ENCODING_LABELS,
    ENCODINGS,
    MIXED_ENCODING,
    SUPPORTED_EXTENSIONS,
    DataLoadError,
    LoadedTable,
    excel_sheets,
    frame_fingerprint,
    guess_gold_columns,
    guess_group_columns,
    guess_id_column,
    guess_text_column,
    list_data_files,
    read_table,
)

# Shaped like the user's gold-labelled sample (a CSV saved as .txt): quoted
# rationales with commas, numeric-looking tier / flag columns, two gold
# columns. The rows themselves are synthetic.
SAMPLE = """\
INDEX,TEXT,GOLD_SENTIMENT,GOLD_INTENT,TIER,CONTESTED,RATIONALE
4000001,Thank you for sorting out my library card renewal so quickly. Brilliant help at the front desk.,positive,praise,1,0,"Explicit thanks, single polarity"
4000002,I want to complain about how I was spoken to at the returns counter on Saturday. It was rude and dismissive.,negative,complaint,1,0,Explicitly framed as a complaint
4000003,I confirm receipt of your email dated 3 May about my reserved items.,neutral,other,1,0,Simple acknowledgement
4000004,Please send me a printed copy of the autumn events programme.,neutral,request,1,0,Simple request for action
4000005,I must say the new opening hours are rather more convenient than I expected.,positive,praise,2,0,Understated praise ('rather more convenient')
4000006,It would be nice if the self-service machines didn't freeze every other day.,negative,complaint,2,0,Complaint phrased as a suggestion
4000007,"Thanks for the quick response; the answer, however, didn't address my question at all.",mixed,complaint,2,1,"Speed praised, substance criticised - arguably negative-dominant"
4000008,"Honestly, after last term's chaos, I expected little - and even that wasn't delivered.",negative,complaint,2,0,"Comparison with the past, clearly negative"
4000009,Marvellous. The 'temporary' closure of the reading room is now in its fourth month.,negative,complaint,3,0,"Sarcasm: positive word, negative meaning"
4000010,Thanks for the letter explaining why my fine can't be waived. Very enlightening.,negative,complaint,3,0,Ironic thanks
4000011,"The branch says the delay is with head office, though I no longer know who to believe.",negative,complaint,3,1,"Distrust or frustration - arguably neutral uncertainty"
4000012,Can a library account be closed on behalf of a relative who has passed away?,neutral,query,3,0,"Sensitive topic, neutral practical question"
"""
SAMPLE_COLUMNS = [
    "INDEX",
    "TEXT",
    "GOLD_SENTIMENT",
    "GOLD_INTENT",
    "TIER",
    "CONTESTED",
    "RATIONALE",
]

# Whatever read_csv(dtype=str) gives on the installed pandas: object on 2.x,
# the str dtype on 3.x. Every format must come back like this.
TEXT_DTYPE = pd.read_csv(io.StringIO("a\nx\n"), dtype=str)["a"].dtype


def write(folder: Path, name: str, content: str | bytes, encoding: str = "utf-8") -> Path:
    path = folder / name
    path.write_bytes(content.encode(encoding) if isinstance(content, str) else content)
    return path


def cells(frame: pd.DataFrame, column: str) -> list[str | None]:
    """Column values with every kind of missing value as None."""
    return [None if pd.isna(value) else value for value in frame[column].tolist()]


def assert_text_frame(frame: pd.DataFrame) -> None:
    assert isinstance(frame.index, pd.RangeIndex)
    assert frame.index.start == 0
    assert all(isinstance(column, str) for column in frame.columns)
    assert frame.columns.is_unique
    for column in frame.columns:
        assert frame[column].dtype == TEXT_DTYPE, column
        for value in frame[column]:
            assert pd.isna(value) or isinstance(value, str), (column, value)


# --- delimited text: encodings ---------------------------------------------------------


def test_reads_plain_utf8_csv(tmp_path):
    path = write(tmp_path, "sample.csv", SAMPLE)
    table = read_table(path)
    assert isinstance(table, LoadedTable)
    assert table.name == "sample.csv"
    assert table.fmt == "csv"
    assert table.encoding == "utf-8"
    assert table.delimiter == ","
    assert table.sheet is None
    assert table.size_bytes == path.stat().st_size
    assert list(table.frame.columns) == SAMPLE_COLUMNS
    assert len(table.frame) == 12
    assert table.frame.loc[0, "RATIONALE"] == "Explicit thanks, single polarity"
    assert table.frame.loc[6, "TEXT"].startswith("Thanks for the quick response; the answer,")
    assert_text_frame(table.frame)


def test_accepts_str_paths(tmp_path):
    path = write(tmp_path, "sample.csv", SAMPLE)
    assert read_table(str(path)).frame.shape == (12, 7)


def test_txt_file_that_is_really_csv(tmp_path):
    # Like the user's sample file: comma-separated, .txt extension.
    table = read_table(write(tmp_path, "text2.txt", SAMPLE))
    assert table.fmt == "csv"
    assert table.delimiter == ","
    assert list(table.frame.columns) == SAMPLE_COLUMNS
    assert cells(table.frame, "TIER")[:3] == ["1", "1", "1"]
    assert table.frame.loc[11, "RATIONALE"] == "Sensitive topic, neutral practical question"


def test_utf8_bom_is_stripped_from_first_column(tmp_path):
    table = read_table(write(tmp_path, "bom.csv", "\ufeffid,text\n1,café\n"))
    assert list(table.frame.columns) == ["id", "text"]
    assert table.encoding == "utf-8-sig"
    assert cells(table.frame, "text") == ["café"]


def test_bom_inside_a_header_name_is_removed(tmp_path):
    table = read_table(write(tmp_path, "odd.csv", "id,\ufefftext \n1,x\n"))
    assert list(table.frame.columns) == ["id", "text"]


def test_cp1252_excel_export(tmp_path):
    # 0x92 = right single quote, 0xA3 = pound sign, 0x96 = en dash in cp1252.
    raw = b"id,text\n1,It\x92s \xa310 \x96 cheap\n2,plain\n"
    table = read_table(write(tmp_path, "excel.csv", raw))
    assert table.encoding == "cp1252"
    assert cells(table.frame, "text") == ["It’s £10 – cheap", "plain"]


def test_latin1_is_the_last_resort(tmp_path):
    # 0x81 is undefined in cp1252, so only latin-1 can decode this.
    table = read_table(write(tmp_path, "weird.csv", b"id,text\n1,a\x81b \xe9\n"))
    assert table.encoding == "latin-1"
    assert cells(table.frame, "text") == ["a\x81b é"]


def test_bom_plus_invalid_utf8_does_not_leak_bom_into_header(tmp_path):
    raw = b"\xef\xbb\xbfid,text\n1,caf\xe9\n"
    table = read_table(write(tmp_path, "mixed.csv", raw))
    assert table.encoding == "cp1252"
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "text") == ["café"]


# What Excel's "Unicode Text (*.txt)" export holds: tab-separated, CRLF,
# text no single-byte code page can represent.
UNICODE_TEXT = "id\ttext\r\n1\tnaïve café – “quoted”, £5\r\n2\t日本語のテキスト\r\n"
UNICODE_CELLS = ["naïve café – “quoted”, £5", "日本語のテキスト"]


def test_utf16_unicode_text_export(tmp_path):
    # Excel writes it as UTF-16 LE with a byte-order mark.
    raw = codecs.BOM_UTF16_LE + UNICODE_TEXT.encode("utf-16-le")
    table = read_table(write(tmp_path, "unicode.txt", raw))
    assert table.encoding == "utf-16-le"
    assert table.delimiter == "\t"
    assert table.fmt == "tsv"
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "id") == ["1", "2"]
    assert cells(table.frame, "text") == UNICODE_CELLS
    assert_text_frame(table.frame)


@pytest.mark.parametrize(
    ("bom", "codec"),
    [
        (codecs.BOM_UTF16_BE, "utf-16-be"),
        (codecs.BOM_UTF32_LE, "utf-32-le"),
        (codecs.BOM_UTF32_BE, "utf-32-be"),
    ],
)
def test_other_unicode_byte_order_marks(tmp_path, bom, codec):
    table = read_table(write(tmp_path, "export.csv", bom + UNICODE_TEXT.encode(codec)))
    assert table.encoding == codec
    assert table.delimiter == "\t"
    assert list(table.frame.columns) == ["id", "text"]  # no BOM left in the header
    assert cells(table.frame, "text") == UNICODE_CELLS


def test_utf32_le_is_not_mistaken_for_utf16_le(tmp_path):
    # The UTF-32 LE byte-order mark starts with the UTF-16 LE one.
    assert codecs.BOM_UTF32_LE.startswith(codecs.BOM_UTF16_LE)
    raw = "id,text\n1,ok\n".encode("utf-32")  # native order, with a BOM
    table = read_table(write(tmp_path, "wide.csv", raw))
    assert table.encoding in ("utf-32-le", "utf-32-be")
    assert cells(table.frame, "text") == ["ok"]


def test_utf8_without_bom_is_still_utf8(tmp_path):
    table = read_table(write(tmp_path, "plain.txt", UNICODE_TEXT))
    assert table.encoding == "utf-8"
    assert table.delimiter == "\t"
    assert cells(table.frame, "text") == UNICODE_CELLS


def test_truncated_utf16_raises(tmp_path):
    raw = codecs.BOM_UTF16_LE + UNICODE_TEXT.encode("utf-16-le")[:-1]
    with pytest.raises(DataLoadError, match=r"cut\.txt as UTF-16 LE"):
        read_table(write(tmp_path, "cut.txt", raw))


def test_jsonl_in_utf16(tmp_path):
    text = '{"id": 1, "text": "naïve – 日本"}\n'
    table = read_table(
        write(tmp_path, "rows.jsonl", codecs.BOM_UTF16_LE + text.encode("utf-16-le"))
    )
    assert table.encoding == "utf-16-le"
    assert cells(table.frame, "text") == ["naïve – 日本"]


@pytest.mark.parametrize("bom", [b"", codecs.BOM_UTF8])
def test_utf8_with_a_few_cp1252_bytes_is_read_as_both(tmp_path, bom):
    # One row pasted from an ANSI file used to turn every other row into
    # mojibake ("TrÃ¨s bien â€” cafÃ©"), since the whole file fell back to cp1252.
    rows = "".join(f"{i},Très bien — café n°{i} ✓\n" for i in range(999))
    raw = bom + ("id,text\n" + rows).encode("utf-8") + "999,naïve “quote”\n".encode("cp1252")
    table = read_table(write(tmp_path, "mixed.csv", raw))
    assert table.encoding == MIXED_ENCODING
    assert ENCODING_LABELS[table.encoding] == "UTF-8 + Windows-1252"
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "text")[0] == "Très bien — café n°0 ✓"
    assert cells(table.frame, "text")[-1] == "naïve “quote”"


def test_mixed_decoding_keeps_bytes_cp1252_leaves_undefined(tmp_path):
    raw = "id,text\n1,café\n2,a".encode() + b"\x81\x92b\n"
    table = read_table(write(tmp_path, "odd.csv", raw))
    assert table.encoding == MIXED_ENCODING
    assert cells(table.frame, "text") == ["café", "a\x81’b"]


def test_every_reported_encoding_has_a_label():
    reported = {"utf-8", *ENCODINGS, "utf-16-le", "utf-16-be", "utf-32-le", "utf-32-be"}
    assert reported | {MIXED_ENCODING} <= set(ENCODING_LABELS)


def test_encoding_and_delimiter_detection_is_logged_at_debug(tmp_path, caplog, monkeypatch):
    # In case an earlier test configured logging with propagation off.
    monkeypatch.setattr(logging.getLogger("batchtextreview"), "propagate", True)
    raw = codecs.BOM_UTF16_LE + UNICODE_TEXT.encode("utf-16-le")
    with caplog.at_level(logging.DEBUG, logger="batchtextreview.data"):
        read_table(write(tmp_path, "unicode.txt", raw))
        read_table(write(tmp_path, "semi.csv", "sep=;\na;b\n1;2\n"))
    records = [r for r in caplog.records if r.name == "batchtextreview.data"]
    assert {r.levelno for r in records} == {logging.DEBUG}
    messages = [r.getMessage() for r in records]
    assert "unicode.txt: byte-order mark found, decoded as utf-16-le" in messages
    assert "unicode.txt: delimiter '\\t' (sniffed)" in messages
    assert "semi.csv: decoded as utf-8" in messages
    assert "semi.csv: delimiter ';' (from its sep= line)" in messages


def test_encodings_constant_order():
    assert ENCODINGS[0] == "utf-8-sig"
    assert ENCODINGS[-1] == "latin-1"


# --- delimited text: delimiters --------------------------------------------------------


@pytest.mark.parametrize(
    ("delimiter", "fmt"),
    [("\t", "tsv"), (";", "csv"), ("|", "csv"), (",", "csv")],
)
@pytest.mark.parametrize("extension", [".csv", ".txt"])
def test_sniffs_delimiter(tmp_path, delimiter, fmt, extension):
    rows = [["id", "text", "score"], ["001", "short", "3,5"], ["002", "a bit longer", "4"]]
    if delimiter == ",":
        rows[1][2] = "3.5"
    content = "\n".join(delimiter.join(row) for row in rows) + "\n"
    table = read_table(write(tmp_path, "data" + extension, content))
    assert table.delimiter == delimiter
    assert table.fmt == fmt
    assert list(table.frame.columns) == ["id", "text", "score"]
    assert cells(table.frame, "id") == ["001", "002"]


def test_semicolon_file_with_unquoted_commas_in_text(tmp_path):
    content = "id;text\n1;Hello, world, again\n2;no commas\n3;one, two\n"
    table = read_table(write(tmp_path, "euro.csv", content))
    assert table.delimiter == ";"
    assert cells(table.frame, "text") == ["Hello, world, again", "no commas", "one, two"]


def test_tsv_prefers_tab_even_with_commas_everywhere(tmp_path):
    content = "id\ttext, with comma\n1\ta, b, c\n2\td, e\n"
    table = read_table(write(tmp_path, "data.tsv", content))
    assert table.delimiter == "\t"
    assert table.fmt == "tsv"
    assert list(table.frame.columns) == ["id", "text, with comma"]


def test_tsv_without_tabs_is_still_sniffed(tmp_path):
    table = read_table(write(tmp_path, "mislabelled.tsv", "id,text\n1,a\n"))
    assert table.delimiter == ","
    assert table.fmt == "csv"


def test_quoted_fields_with_commas_newlines_and_quotes(tmp_path):
    content = 'id,text,label\n1,"Line one\nline two, with comma",pos\n2,"He said ""hi""",neg\n3,plain,neu\n'
    table = read_table(write(tmp_path, "quoted.csv", content))
    assert cells(table.frame, "text") == [
        "Line one\nline two, with comma",
        'He said "hi"',
        "plain",
    ]
    assert cells(table.frame, "label") == ["pos", "neg", "neu"]


def test_windows_line_endings(tmp_path):
    table = read_table(write(tmp_path, "crlf.csv", "id,text\r\n1,a\r\n2,b\r\n"))
    assert cells(table.frame, "text") == ["a", "b"]
    assert list(table.frame.columns) == ["id", "text"]


def test_single_column_file_with_commas_in_text(tmp_path):
    content = "comment\nGreat service, thanks\nToo slow, too expensive, too rude\nFine\n"
    table = read_table(write(tmp_path, "comments.txt", content))
    assert list(table.frame.columns) == ["comment"]
    assert cells(table.frame, "comment") == [
        "Great service, thanks",
        "Too slow, too expensive, too rude",
        "Fine",
    ]
    assert table.delimiter == ","


def test_single_column_quoted_csv(tmp_path):
    table = read_table(write(tmp_path, "one.csv", 'text\n"a, b"\n"multi\nline"\n'))
    assert cells(table.frame, "text") == ["a, b", "multi\nline"]


@pytest.mark.parametrize(
    ("content", "expected"),
    [
        ('text\n"Best" product ever\nok then\n', ['"Best" product ever', "ok then"]),
        (
            'text\nline one\n"quoted start no end\nline three\n',
            ["line one", '"quoted start no end', "line three"],
        ),
        ('text\n"I never said that\nsecond\n', ['"I never said that', "second"]),
    ],
)
def test_single_column_text_keeps_quote_marks_as_written(tmp_path, content, expected):
    table = read_table(write(tmp_path, "lines.txt", content))
    assert list(table.frame.columns) == ["text"]
    assert cells(table.frame, "text") == expected


def test_sniffing_only_looks_at_the_start_of_big_files(tmp_path):
    row = "1;" + "x" * 200 + "\n"
    content = "id;text\n" + row * 2000  # ~400 KB
    table = read_table(write(tmp_path, "big.csv", content))
    assert table.delimiter == ";"
    assert len(table.frame) == 2000


TABLE_ROWS = '"1","Parcel late, box damaged","negative"\r\n"2","Great","positive"\r\n'


@pytest.mark.parametrize(
    ("content", "columns"),
    [
        ('"Id","Text","Label"\r\n' + TABLE_ROWS, ["Id", "Text", "Label"]),
        ('"Id","Text"\r\n"1","Parcel late, box damaged"\r\n"2","Great"\r\n', ["Id", "Text"]),
    ],
)
def test_powershell_type_line_is_skipped(tmp_path, content, columns):
    # Windows PowerShell 5.1's Export-Csv writes this line unless told not to;
    # it used to turn the file into one column of raw CSV lines.
    content = "#TYPE System.Management.Automation.PSCustomObject\r\n" + content
    table = read_table(write(tmp_path, "export.csv", content))
    assert list(table.frame.columns) == columns
    assert cells(table.frame, "Text") == ["Parcel late, box damaged", "Great"]


@pytest.mark.parametrize(
    "preamble",
    [
        "Customer feedback export - Q3 2025\r\n",
        "Customer feedback export\r\n\r\nGenerated 01/10/2025 09:00\r\n",
        '"Quoted report title"\n',
    ],
)
def test_title_lines_above_the_table_are_skipped(tmp_path, preamble):
    content = preamble + '"Id","Text","Label"\r\n' + TABLE_ROWS + '"3","Fine","neutral"\r\n'
    table = read_table(write(tmp_path, "report.csv", content))
    assert list(table.frame.columns) == ["Id", "Text", "Label"]
    assert cells(table.frame, "Text") == ["Parcel late, box damaged", "Great", "Fine"]
    assert guess_text_column(table.frame) == "Text"


def test_title_line_keeps_line_numbers_in_errors(tmp_path):
    content = "Report\nId,Text,Label\n" + "1,a,x\n" * 30 + "2,b,y,extra\n" + "3,c,z\n" * 3
    with pytest.raises(DataLoadError, match="line 33 has 4 fields"):
        read_table(write(tmp_path, "report.csv", content))


@pytest.mark.parametrize(
    "content",
    [
        "name\nSmith, John\nDoe, Jane\nRoe, Rick\nLee, Ann\n",  # two "columns": too risky
        "comment\nRed, green, blue\nOne, two, three\n",  # only one row below
        "comment\nRed, green, blue\nOne, two, three\nFour, five\n",  # not all alike
    ],
)
def test_single_column_lines_are_not_mistaken_for_a_titled_table(tmp_path, content):
    table = read_table(write(tmp_path, "list.csv", content))
    assert table.frame.shape[1] == 1
    assert cells(table.frame, table.frame.columns[0]) == content.splitlines()[1:]


@pytest.mark.parametrize(
    "content",
    [
        # The first line has a comma, most lines don't.
        "Customer comments (free text, optional)\nArrived broken\nOkay, I guess not\n"
        "Never again\nGood\n",
        "Great product, fast shipping\nArrived broken\nOkay, I guess\nNever again\nGood\n",
        # A later line has more commas than the first.
        "Hello, this is the first note.\nSecond note\nThird, note, with, many, commas\n",
        # Every line has commas, but always followed by a space, as in prose.
        "Review (verified buyer, 2024)\nI loved it, the colour is great, and it came fast.\n"
        "Good value, would buy again.\nTerrible, broke after a week, never again, avoid.\n",
        "Customer comments (free text, optional)\nArrived broken, box crushed\n"
        "Okay, I guess not\nNever again\nGood, fast\n",
    ],
)
def test_one_text_per_line_with_commas_is_not_split(tmp_path, content):
    # These used to be cut at the first comma ("Okay, I guess not" -> "Okay")
    # or fail with advice to check the commas.
    table = read_table(write(tmp_path, "comments.txt", content))
    lines = content.splitlines()
    assert list(table.frame.columns) == [lines[0]]
    assert cells(table.frame, lines[0]) == lines[1:]


def test_ragged_table_is_still_a_table(tmp_path):
    table = read_table(write(tmp_path, "ragged.csv", "id,text,label\n1,a\n2,b\n3,c,pos\n"))
    assert list(table.frame.columns) == ["id", "text", "label"]
    assert cells(table.frame, "label") == [None, None, "pos"]
    # Hand-typed, with no trailing comma on the rows that only have an ID.
    table = read_table(write(tmp_path, "typed.csv", "id,comment\n1\n2,nice one\n3\n4\n"))
    assert list(table.frame.columns) == ["id", "comment"]
    assert cells(table.frame, "comment") == [None, "nice one", None, None]


def test_small_table_with_an_unquoted_comma_still_explains(tmp_path):
    with pytest.raises(DataLoadError, match=r"line 2 has 4 fields but the header has 3"):
        read_table(write(tmp_path, "broken.csv", "id,text,label\n1,hello, world,pos\n2,b,neg\n"))


def test_malformed_row_raises_with_line_and_hint(tmp_path):
    path = write(tmp_path, "bad.csv", "id,text\n1,ok\n2,unquoted, comma\n")
    with pytest.raises(DataLoadError, match=r"bad\.csv") as info:
        read_table(path)
    message = str(info.value)
    assert "line 3 has 3 fields but the header has 2" in message
    assert "quotes" in message
    assert "C error" not in message


def _numbered_rows(count: int) -> list[str]:
    return ["id,text"] + [f"{number},row {number}" for number in range(1, count + 1)]


@pytest.mark.parametrize(
    ("bad_row", "line"),
    [(50, 51), (100, 101)],  # middle of the file, last row
)
def test_unclosed_quote_raises_instead_of_dropping_rows(tmp_path, bad_row, line):
    rows = _numbered_rows(100)
    rows[bad_row] = f'{bad_row},"unclosed quote here'
    path = write(tmp_path, "quote.csv", "\n".join(rows) + "\n")
    with pytest.raises(DataLoadError, match=r"quote\.csv") as info:
        read_table(path)
    message = str(info.value)
    assert f"row starting on line {line} is never closed" in message
    assert '("")' in message
    assert "EOF" not in message


def test_unclosed_quote_after_a_ragged_row_still_raises(tmp_path):
    content = 'id,text\n1,a,,\n2,"never closed\n3,c\n4,d\n'
    with pytest.raises(DataLoadError, match="row starting on line 3 is never closed"):
        read_table(write(tmp_path, "quote.csv", content))


def test_unclosed_quote_in_a_big_file_names_its_line(tmp_path):
    # The unclosed field swallows > 128 KB, past csv's field size limit.
    rows = [f"{row},{'x' * 60}" for row in _numbered_rows(5000)]
    rows[50] = '50,"unclosed quote here'
    with pytest.raises(DataLoadError, match="row starting on line 51 is never closed"):
        read_table(write(tmp_path, "big.csv", "\n".join(rows) + "\n"))


def test_ragged_rows_keep_every_row_around_odd_quotes(tmp_path):
    # A trailing-comma row used to send the file through pandas' python
    # parser, which silently dropped the '"Best" product' row.
    content = 'id,text\n1,a,,\n2,"Best" product\n3,c\n4,d,\n'
    table = read_table(write(tmp_path, "ragged.csv", content))
    assert cells(table.frame, "id") == ["1", "2", "3", "4"]
    assert cells(table.frame, "text") == ["a", "Best product", "c", "d"]


def test_rows_of_growing_width_with_empty_extras_are_trimmed(tmp_path):
    content = "id,text\n1,a,\n2,b,,\n3,c,,,\n4,d,,,,\n"
    table = read_table(write(tmp_path, "ragged.csv", content))
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "text") == ["a", "b", "c", "d"]


# --- delimited text: values and columns ------------------------------------------------


def test_na_like_text_is_kept_and_only_empty_cells_are_missing(tmp_path):
    content = 'id,text\n1,NA\n2,None\n3,null\n4,NaN\n5,n/a\n6,\n7,""\n8, \n'
    table = read_table(write(tmp_path, "na.csv", content))
    values = cells(table.frame, "text")
    assert values[:5] == ["NA", "None", "null", "NaN", "n/a"]
    assert values[5] is None
    assert values[6] is None
    assert values[7] == " "
    assert table.frame["text"].isna().tolist() == [False] * 5 + [True, True, False]


def test_numbers_stay_as_written(tmp_path):
    content = "id,score,code,big\n007,1.50,1e3,12345678901234567890\n010,2,0x1F,-0\n"
    table = read_table(write(tmp_path, "nums.csv", content))
    assert cells(table.frame, "id") == ["007", "010"]
    assert cells(table.frame, "score") == ["1.50", "2"]
    assert cells(table.frame, "code") == ["1e3", "0x1F"]
    assert cells(table.frame, "big") == ["12345678901234567890", "-0"]
    assert_text_frame(table.frame)


def test_duplicate_column_names_get_suffixes(tmp_path):
    table = read_table(write(tmp_path, "dups.csv", "a,b,a,a\n1,2,3,4\n"))
    assert list(table.frame.columns) == ["a", "b", "a (2)", "a (3)"]
    assert cells(table.frame, "a (3)") == ["4"]


def test_duplicate_suffix_never_clobbers_a_real_column(tmp_path):
    table = read_table(write(tmp_path, "dups.csv", "a,a,a (2)\n1,2,3\n"))
    assert list(table.frame.columns) == ["a", "a (3)", "a (2)"]
    assert cells(table.frame, "a (2)") == ["3"]


def test_header_names_are_stripped(tmp_path):
    table = read_table(write(tmp_path, "ws.csv", "  id , text\t\n1,x\n"))
    assert list(table.frame.columns) == ["id", "text"]


@pytest.mark.parametrize("first_line", ["\t", "   ", " \t ", "\n  "])
def test_whitespace_lines_before_the_header_are_skipped(tmp_path, first_line):
    content = f"{first_line}\nid,text\n1,a\n"
    table = read_table(write(tmp_path, "lead.csv", content))
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "text") == ["a"]
    assert table.delimiter == ","


def test_line_numbers_count_skipped_leading_lines(tmp_path):
    content = "  \n\nid,text\n1,a\n2,b,c\n"
    with pytest.raises(DataLoadError, match="line 5 has 3 fields"):
        read_table(write(tmp_path, "lead.csv", content))


@pytest.mark.parametrize(
    ("content", "delimiter"),
    [
        ("sep=;\nid;text\n1;a, b\n", ";"),
        ("SEP=|\r\nid|text\r\n1|a, b\r\n", "|"),
        ("sep=\t\nid\ttext\n1\ta, b\n", "\t"),
    ],
)
def test_excel_sep_line_sets_the_delimiter(tmp_path, content, delimiter):
    table = read_table(write(tmp_path, "hinted.csv", content))
    assert table.delimiter == delimiter
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "text") == ["a, b"]


def test_empty_unnamed_columns_are_dropped(tmp_path):
    # Trailing delimiters, as Excel writes when a stray cell was once used.
    content = "id,text,,\n1,a,,\n2,b,,\n"
    table = read_table(write(tmp_path, "trailing.csv", content))
    assert list(table.frame.columns) == ["id", "text"]


def test_unnamed_column_with_values_is_kept(tmp_path):
    # A pandas to_csv() index column has an empty header but real values.
    content = ",text,empty_named\n0,a,\n1,b,\n"
    table = read_table(write(tmp_path, "indexed.csv", content))
    assert list(table.frame.columns) == ["Unnamed: 0", "text", "empty_named"]
    assert cells(table.frame, "Unnamed: 0") == ["0", "1"]
    assert table.frame["empty_named"].isna().all()


def test_fully_empty_rows_are_dropped(tmp_path):
    content = "id,text\n1,a\n,\n\n2,b\n,,\n"
    table = read_table(write(tmp_path, "gaps.csv", content))
    assert cells(table.frame, "id") == ["1", "2"]
    assert list(table.frame.index) == [0, 1]


def test_rows_with_extra_empty_fields_are_trimmed(tmp_path):
    content = 'id,text\n1,a,,\n2,"b, quoted",\n3,c\n'
    table = read_table(write(tmp_path, "ragged.csv", content))
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "text") == ["a", "b, quoted", "c"]


def test_rows_with_extra_values_still_fail(tmp_path):
    content = 'id,text\n1,a,,\n2,"multi\nline",\n3,b,surprise\n'
    with pytest.raises(DataLoadError, match="line 5 has 3 fields but the header has 2"):
        read_table(write(tmp_path, "ragged.csv", content))


def test_short_rows_are_padded_with_missing(tmp_path):
    table = read_table(write(tmp_path, "short.csv", "id,text,label\n1,a\n2,b,x\n"))
    assert cells(table.frame, "label") == [None, "x"]


@pytest.mark.parametrize(
    ("content", "pattern"),
    [
        (b"", "empty"),
        (b"   \n\n  \r\n", "empty"),
        (b"id,text\n", "No data rows"),
        (b"id,text\n,\n,\n", "No data rows"),
        (b",,\n", "No columns"),
    ],
)
def test_empty_tables_raise(tmp_path, content, pattern):
    with pytest.raises(DataLoadError, match=pattern):
        read_table(write(tmp_path, "empty.csv", content))


def test_binary_data_in_a_text_file_raises(tmp_path):
    with pytest.raises(DataLoadError, match="binary"):
        read_table(write(tmp_path, "blob.csv", b"id,text\n1,\x00\x01\x02\n"))


def test_workbook_renamed_to_csv_gets_a_helpful_error(tmp_path):
    xlsx = tmp_path / "real.xlsx"
    pd.DataFrame({"a": ["1"]}).to_excel(xlsx, index=False)
    path = write(tmp_path, "renamed.csv", xlsx.read_bytes())
    with pytest.raises(DataLoadError, match=r"looks like an Excel workbook.*\.xlsx"):
        read_table(path)


def test_parquet_renamed_to_csv_gets_a_helpful_error(tmp_path):
    source = tmp_path / "real.parquet"
    pd.DataFrame({"a": [1]}).to_parquet(source)
    with pytest.raises(DataLoadError, match="Parquet"):
        read_table(write(tmp_path, "renamed.txt", source.read_bytes()))


# --- Excel ------------------------------------------------------------------------------


@pytest.fixture
def workbook(tmp_path) -> Path:
    path = tmp_path / "book.xlsx"
    first = pd.DataFrame(
        {
            "id": ["007", "008", "009"],
            "text": ["Hello, world", "NA", None],
            "score": [1, 2.5, None],
        }
    )
    second = pd.DataFrame({"comment": ["x", "y"], "gold": ["pos", "neg"]})
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        first.to_excel(writer, sheet_name="Main", index=False)
        second.to_excel(writer, sheet_name="Other data", index=False)
    return path


def test_reads_first_sheet_by_default(workbook):
    table = read_table(workbook)
    assert table.fmt == "excel"
    assert table.sheet == "Main"
    assert table.encoding is None
    assert table.delimiter is None
    assert list(table.frame.columns) == ["id", "text", "score"]
    assert cells(table.frame, "id") == ["007", "008", "009"]
    assert cells(table.frame, "text") == ["Hello, world", "NA", None]
    assert cells(table.frame, "score") == ["1", "2.5", None]
    assert_text_frame(table.frame)


def test_reads_chosen_sheet(workbook):
    table = read_table(workbook, sheet="Other data")
    assert table.sheet == "Other data"
    assert list(table.frame.columns) == ["comment", "gold"]
    assert cells(table.frame, "gold") == ["pos", "neg"]


def test_missing_sheet_lists_the_real_ones(workbook):
    with pytest.raises(DataLoadError, match=r"'Nope' not found.*Main, Other data"):
        read_table(workbook, sheet="Nope")


def test_excel_sheets_from_path_bytes_and_file(workbook):
    expected = ["Main", "Other data"]
    assert excel_sheets(workbook) == expected
    assert excel_sheets(str(workbook)) == expected
    assert excel_sheets(workbook.read_bytes()) == expected
    assert excel_sheets(io.BytesIO(workbook.read_bytes())) == expected


def test_excel_sheets_is_empty_for_anything_unreadable(tmp_path):
    assert excel_sheets(write(tmp_path, "a.csv", "a,b\n1,2\n")) == []
    assert excel_sheets(b"not a workbook") == []
    assert excel_sheets(tmp_path / "missing.xlsx") == []
    assert excel_sheets(tmp_path) == []


def test_xlsm_extension_reads_as_excel(workbook, tmp_path):
    path = write(tmp_path, "macro.xlsm", workbook.read_bytes())
    assert read_table(path).fmt == "excel"


def test_excel_duplicate_and_blank_headers(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["text", "text", None, " label ", None])
    sheet.append(["a", "b", None, "x", None])
    sheet.append(["c", "d", "stray", "y", None])
    sheet.append([None, None, None, None, None])
    path = tmp_path / "messy.xlsx"
    book.save(path)
    table = read_table(path)
    assert list(table.frame.columns) == ["text", "text (2)", "Unnamed: 2", "label"]
    assert cells(table.frame, "Unnamed: 2") == [None, "stray"]
    assert len(table.frame) == 2


def test_empty_sheet_raises(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    book.active.title = "Blank"
    path = tmp_path / "blank.xlsx"
    book.save(path)
    with pytest.raises(DataLoadError, match="empty"):
        read_table(path)


def save_sheet(tmp_path: Path, rows: list[list[str | None]], *, first_row: int = 1,
               first_column: int = 1, merge: str | None = None) -> Path:  # fmt: skip
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    for r, row in enumerate(rows, start=first_row):
        for c, value in enumerate(row, start=first_column):
            if value is not None:
                sheet.cell(r, c, value)
    if merge:
        sheet.merge_cells(merge)
    path = tmp_path / "sheet.xlsx"
    book.save(path)
    return path


FEEDBACK_ROWS = [
    ["ID", "Text", "Label"],
    ["1", "Parcel late", "negative"],
    ["2", "Lovely staff", "positive"],
]


def test_excel_table_below_empty_rows_and_columns(tmp_path):
    # Used to give "Unnamed: 1/2/3" columns with the header row as data.
    table = read_table(save_sheet(tmp_path, FEEDBACK_ROWS, first_row=3, first_column=2))
    assert list(table.frame.columns) == ["ID", "Text", "Label"]
    assert cells(table.frame, "Text") == ["Parcel late", "Lovely staff"]


@pytest.mark.parametrize("merge", ["A1:C1", None])
def test_excel_title_row_above_the_header_is_skipped(tmp_path, merge):
    rows = [["Customer feedback Q3 2025"], [None], *FEEDBACK_ROWS]
    table = read_table(save_sheet(tmp_path, rows, merge=merge))
    assert list(table.frame.columns) == ["ID", "Text", "Label"]
    assert cells(table.frame, "ID") == ["1", "2"]
    assert guess_text_column(table.frame) == "Text"
    assert guess_gold_columns(table.frame, exclude=("Text",))[0] == "Label"


@pytest.mark.parametrize(
    ("rows", "columns"),
    [
        # A one-column list with a stray wide row is not a title + table.
        ([["Text"], ["a"], ["b"], ["c", "d", "e"], ["f"]], ["Text", "Unnamed: 1", "Unnamed: 2"]),
        # A header with one name over two filled columns.
        ([["comment", None], ["hello", "x"], ["bye", "y"]], ["comment", "Unnamed: 1"]),
        # Header names that repeat in their own column are data.
        ([["Text"], ["a", "x", "y"], ["b", "x", "z"]], ["Text", "Unnamed: 1", "Unnamed: 2"]),
    ],
)
def test_excel_single_cell_first_row_is_kept_when_no_table_follows(tmp_path, rows, columns):
    table = read_table(save_sheet(tmp_path, rows))
    assert list(table.frame.columns) == columns
    assert len(table.frame) == len(rows) - 1


def test_hidden_first_sheet_is_not_the_default(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    lists = book.active
    lists.title = "Lists"
    for value in ("positive", "negative", "neutral"):
        lists.append([value])
    lists.sheet_state = "hidden"
    tickets = book.create_sheet("Tickets")
    for row in FEEDBACK_ROWS:
        tickets.append(row)
    book.create_sheet("Chart data").append(["x", "y"])
    path = tmp_path / "lists.xlsx"
    book.save(path)
    table = read_table(path)
    assert table.sheet == "Tickets"
    assert list(table.frame.columns) == ["ID", "Text", "Label"]
    # The UI offers the default first; the hidden sheet stays reachable.
    assert excel_sheets(path) == ["Tickets", "Chart data", "Lists"]
    assert read_table(path, sheet="Lists").sheet == "Lists"


def test_corrupt_workbook_raises(tmp_path):
    path = write(tmp_path, "broken.xlsx", b"PK\x03\x04 this is not really a zip")
    with pytest.raises(DataLoadError, match=r"Couldn't read broken\.xlsx as an Excel workbook"):
        read_table(path)


# --- Parquet ----------------------------------------------------------------------------


def test_reads_parquet_as_text(tmp_path):
    path = tmp_path / "data.parquet"
    pd.DataFrame(
        {
            "id": pd.array([1, 2, None], dtype="Int64"),
            "ratio": [0.5, 2.0, None],
            "flag": [True, False, True],
            "text": ["NA", None, "None"],
        }
    ).to_parquet(path)
    table = read_table(path)
    assert table.fmt == "parquet"
    assert table.encoding is None
    assert cells(table.frame, "id") == ["1", "2", None]
    assert cells(table.frame, "ratio") == ["0.5", "2.0", None]
    assert cells(table.frame, "flag") == ["True", "False", "True"]
    assert cells(table.frame, "text") == ["NA", None, "None"]
    assert_text_frame(table.frame)


def test_parquet_int_column_with_gaps_reads_without_decimal(tmp_path):
    path = tmp_path / "gappy.parquet"
    # Written from plain numpy ints with a gap: pandas would read it back as
    # float64 ("3.0") without the nullable backend.
    import pyarrow as pa
    import pyarrow.parquet as pq

    columns = {"text": ["a", "b", "c"], "label": pa.array([3, None, 1], type=pa.int64())}
    pq.write_table(pa.table(columns), path)
    assert cells(read_table(path).frame, "label") == ["3", None, "1"]


def test_parquet_named_index_becomes_a_column(tmp_path):
    path = tmp_path / "indexed.parquet"
    frame = pd.DataFrame({"case_id": ["a", "b"], "text": ["x", "y"]}).set_index("case_id")
    frame.to_parquet(path)
    table = read_table(path)
    assert list(table.frame.columns) == ["case_id", "text"]
    assert isinstance(table.frame.index, pd.RangeIndex)


def test_parquet_index_that_copies_a_column_is_not_duplicated(tmp_path):
    path = tmp_path / "kept.parquet"
    frame = pd.DataFrame({"id": ["a1", "a2"], "text": ["x", "y"]})
    frame.set_index("id", drop=False).to_parquet(path)
    table = read_table(path)
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "id") == ["a1", "a2"]


def test_parquet_index_named_like_a_different_column_is_kept(tmp_path):
    path = tmp_path / "clash.parquet"
    index = pd.Index(["z1", "z2"], name="id")
    pd.DataFrame({"id": ["a1", "a2"], "text": ["x", "y"]}, index=index).to_parquet(path)
    table = read_table(path)
    assert list(table.frame.columns) == ["id", "id (2)", "text"]
    assert cells(table.frame, "id") == ["z1", "z2"]
    assert cells(table.frame, "id (2)") == ["a1", "a2"]


def test_parquet_unnamed_index_is_dropped(tmp_path):
    path = tmp_path / "filtered.parquet"
    frame = pd.DataFrame({"text": ["x", "y", "z"]}, index=[10, 20, 30])
    frame.to_parquet(path)
    table = read_table(path)
    assert list(table.frame.columns) == ["text"]
    assert list(table.frame.index) == [0, 1, 2]


def test_corrupt_parquet_raises(tmp_path):
    with pytest.raises(DataLoadError, match=r"Couldn't read junk\.parquet as a Parquet file"):
        read_table(write(tmp_path, "junk.parquet", b"definitely not parquet"))


# --- JSON -------------------------------------------------------------------------------


def test_reads_jsonl(tmp_path):
    lines = [
        {"id": "007", "text": "NA", "score": 3, "ratio": 1.5, "ok": True, "tags": ["a", "b"]},
        {"id": "008", "text": None, "score": None},
        {"id": "009", "text": "café", "extra": {"k": 1}},
    ]
    content = "\n".join(json.dumps(line, ensure_ascii=False) for line in lines) + "\n\n"
    table = read_table(write(tmp_path, "data.jsonl", content))
    assert table.fmt == "jsonl"
    assert table.encoding == "utf-8"
    frame = table.frame
    assert list(frame.columns) == ["id", "text", "score", "ratio", "ok", "tags", "extra"]
    assert cells(frame, "id") == ["007", "008", "009"]
    assert cells(frame, "text") == ["NA", None, "café"]
    assert cells(frame, "score") == ["3", None, None]
    assert cells(frame, "ratio") == ["1.5", None, None]
    assert cells(frame, "ok") == ["true", None, None]
    assert cells(frame, "tags") == ['["a", "b"]', None, None]
    assert cells(frame, "extra") == [None, None, '{"k": 1}']
    assert_text_frame(frame)


def test_jsonl_bad_line_reports_line_number(tmp_path):
    content = '{"a": 1}\n\n{"a": 2\n'
    with pytest.raises(DataLoadError, match=r"line 3 of data\.jsonl"):
        read_table(write(tmp_path, "data.jsonl", content))


def test_jsonl_error_says_where_on_the_line(tmp_path):
    with pytest.raises(DataLoadError, match="Unterminated string starting at column 7"):
        read_table(write(tmp_path, "data.jsonl", '{"a": "no end\n'))


@pytest.mark.parametrize("extension", [".jsonl", ".json"])
def test_jsonl_keeps_unicode_line_breaks_inside_strings(tmp_path, extension):
    texts = ["line one\u2028still one", "para\u0085graph", "x\u2029y\x0cz"]
    content = pd.DataFrame({"id": [1, 2, 3], "text": texts}).to_json(
        orient="records", lines=True, force_ascii=False
    )
    table = read_table(write(tmp_path, f"data{extension}", content))
    assert table.fmt == "jsonl"
    assert cells(table.frame, "text") == texts


def test_jsonl_line_that_is_not_an_object(tmp_path):
    with pytest.raises(DataLoadError, match="Line 2"):
        read_table(write(tmp_path, "data.jsonl", '{"a": 1}\n[1, 2]\n'))


def test_reads_json_records(tmp_path):
    records = [{"id": 1, "text": "a"}, {"id": 2, "text": "b", "label": "pos"}]
    table = read_table(write(tmp_path, "data.json", json.dumps(records)))
    assert table.fmt == "json"
    assert list(table.frame.columns) == ["id", "text", "label"]
    assert cells(table.frame, "id") == ["1", "2"]
    assert cells(table.frame, "label") == [None, "pos"]
    assert_text_frame(table.frame)


@pytest.mark.parametrize("orient", ["records", "columns", "split", "table", "index_wrapped"])
def test_reads_common_json_layouts(tmp_path, orient):
    frame = pd.DataFrame({"id": ["001", "002"], "text": ["hello", "NA"]})
    if orient == "index_wrapped":
        content = json.dumps({"items": frame.to_dict(orient="records")})
    else:
        content = frame.to_json(orient=orient)
    table = read_table(write(tmp_path, "data.json", content))
    assert list(table.frame.columns) == ["id", "text"]
    assert cells(table.frame, "id") == ["001", "002"]
    assert cells(table.frame, "text") == ["hello", "NA"]


def test_json_columns_layout_lines_rows_up_by_key(tmp_path):
    content = json.dumps(
        {
            "id": {"0": "a", "1": "b", "2": "c"},
            "text": {"1": "TEXT-B", "0": "TEXT-A"},  # out of order, "2" missing
        }
    )
    table = read_table(write(tmp_path, "data.json", content))
    assert cells(table.frame, "id") == ["a", "b", "c"]
    assert cells(table.frame, "text") == ["TEXT-A", "TEXT-B", None]


def test_json_lines_saved_as_json(tmp_path):
    content = '{"text": "a"}\n{"text": "b"}\n'
    table = read_table(write(tmp_path, "data.json", content))
    assert table.fmt == "jsonl"
    assert cells(table.frame, "text") == ["a", "b"]


def test_json_with_bom(tmp_path):
    table = read_table(write(tmp_path, "bom.json", '\ufeff[{"text": "x"}]'))
    assert list(table.frame.columns) == ["text"]


@pytest.mark.parametrize(
    ("content", "pattern"),
    [
        ("{not json", "Couldn't parse data.json as JSON"),
        ("42", "array of objects"),
        ('["a", "b"]', "array of objects"),
        ('{"a": 1, "b": 2}', "array of objects"),
        ("[]", "No columns"),
    ],
)
def test_json_that_is_not_a_table_raises(tmp_path, content, pattern):
    with pytest.raises(DataLoadError, match=re.escape(pattern)):
        read_table(write(tmp_path, "data.json", content))


# --- sources -----------------------------------------------------------------------------


def test_bytes_source_needs_a_name():
    with pytest.raises(DataLoadError, match="file name"):
        read_table(SAMPLE.encode())


def test_bytes_source_with_name():
    table = read_table(SAMPLE.encode(), name="upload.csv")
    assert table.name == "upload.csv"
    assert table.size_bytes == len(SAMPLE.encode())
    assert list(table.frame.columns) == SAMPLE_COLUMNS


def test_bytearray_and_memoryview_sources():
    raw = b"id,text\n1,a\n"
    assert read_table(bytearray(raw), name="a.csv").frame.shape == (1, 2)
    assert read_table(memoryview(raw), name="a.csv").frame.shape == (1, 2)


def test_name_only_keeps_the_basename():
    table = read_table(b"id,text\n1,a\n", name="C:/Users/me/data/upload.csv")
    assert table.name == "upload.csv"


def test_binary_file_object_uses_its_name(tmp_path):
    path = write(tmp_path, "opened.tsv", "id\ttext\n1\ta\n")
    with path.open("rb") as handle:
        table = read_table(handle)
    assert table.name == "opened.tsv"
    assert table.fmt == "tsv"


def test_file_object_without_name_needs_one():
    with pytest.raises(DataLoadError, match="file name"):
        read_table(io.BytesIO(b"id,text\n1,a\n"))
    table = read_table(io.BytesIO(b"id,text\n1,a\n"), name="x.csv")
    assert cells(table.frame, "text") == ["a"]


class FakeUpload(io.BytesIO):
    """Mimics Streamlit's UploadedFile: a BytesIO with a .name."""

    def __init__(self, data: bytes, name: str) -> None:
        super().__init__(data)
        self.name = name


def test_upload_object_can_be_read_on_every_rerun():
    upload = FakeUpload(SAMPLE.encode(), "text2.txt")
    first = read_table(upload)
    upload.read()  # someone else consumed it
    second = read_table(upload)
    assert first.name == "text2.txt"
    pd.testing.assert_frame_equal(first.frame, second.frame)


def test_seekable_stream_without_getvalue_is_read_from_the_start(tmp_path):
    path = write(tmp_path, "stream.csv", "id,text\n1,a\n")
    with path.open("rb") as handle:
        handle.read(3)
        table = read_table(io.BufferedReader(handle), name="stream.csv")
    assert list(table.frame.columns) == ["id", "text"]


def test_explicit_name_overrides_path_extension(tmp_path):
    path = write(tmp_path, "export.dat", "id;text\n1;a\n")
    table = read_table(path, name="export.csv")
    assert table.delimiter == ";"
    assert table.name == "export.csv"


def test_excel_from_bytes_and_upload(workbook):
    raw = workbook.read_bytes()
    assert read_table(raw, name="book.xlsx", sheet="Other data").sheet == "Other data"
    assert read_table(FakeUpload(raw, "book.xlsx")).sheet == "Main"


@pytest.mark.parametrize(
    ("name", "pattern"),
    [
        ("report.docx", r"Unsupported file type '\.docx'.*\.csv, \.tsv"),
        ("old.xls", r"legacy \.xls.*save it as \.xlsx"),
        ("README", "no file extension"),
    ],
)
def test_unsupported_formats(tmp_path, name, pattern):
    with pytest.raises(DataLoadError, match=pattern):
        read_table(write(tmp_path, name, b"id,text\n1,a\n"))


def test_extension_is_case_insensitive(tmp_path):
    assert read_table(write(tmp_path, "LOUD.CSV", "id,text\n1,a\n")).fmt == "csv"


def test_missing_file_and_folder(tmp_path):
    with pytest.raises(DataLoadError, match="File not found"):
        read_table(tmp_path / "nope.csv")
    with pytest.raises(DataLoadError, match="folder"):
        read_table(tmp_path)


def test_data_load_error_is_a_value_error():
    assert issubclass(DataLoadError, ValueError)


def test_dtype_is_the_same_for_every_format(tmp_path, workbook):
    frame = pd.DataFrame({"id": ["1", "2"], "text": ["a", None]})
    frame.to_parquet(tmp_path / "f.parquet")
    paths = [
        write(tmp_path, "f.csv", "id,text\n1,a\n2,\n"),
        write(tmp_path, "f.json", frame.to_json(orient="records")),
        write(tmp_path, "f.jsonl", frame.to_json(orient="records", lines=True)),
        tmp_path / "f.parquet",
        workbook,
    ]
    for path in paths:
        assert_text_frame(read_table(path).frame)


def test_supported_extensions_are_lowercase_dotted():
    assert all(ext.startswith(".") and ext == ext.lower() for ext in SUPPORTED_EXTENSIONS)


# --- guessing ----------------------------------------------------------------------------


@pytest.fixture
def sample() -> pd.DataFrame:
    return read_table(SAMPLE.encode(), name="text2.txt").frame


def test_guesses_on_the_sample_data(sample):
    assert guess_text_column(sample) == "TEXT"
    assert guess_id_column(sample) == "INDEX"
    assert guess_gold_columns(sample, exclude=("TEXT", "INDEX")) == [
        "GOLD_SENTIMENT",
        "GOLD_INTENT",
        "TIER",
        "CONTESTED",
    ]
    assert guess_group_columns(sample, exclude=("GOLD_SENTIMENT",)) == [
        "GOLD_INTENT",
        "TIER",
        "CONTESTED",
    ]


def test_text_column_name_priority_and_case():
    frame = pd.DataFrame({"Notes": ["n"], "BODY": ["b"], "comment": ["c"]})
    assert guess_text_column(frame) == "BODY"
    assert guess_text_column(pd.DataFrame({" Review ": ["short"], "x": ["much longer"]})) == (
        " Review "
    )


def test_text_column_falls_back_to_longest_mean_length():
    frame = pd.DataFrame(
        {
            "id": ["1", "2", "3"],
            "feedback": ["a long piece of feedback", "another one", "and one more"],
            "label": ["pos", "neg", "pos"],
        }
    )
    assert guess_text_column(frame) == "feedback"


def test_text_column_tie_goes_left():
    frame = pd.DataFrame({"left": ["abc", "de"], "right": ["de", "abc"]})
    assert guess_text_column(frame) == "left"


def test_text_column_counts_missing_as_short():
    frame = pd.DataFrame(
        {
            "remark": ["a really long remark that only appears once", None, None, None],
            "feedback": ["medium length text", "medium length text", "medium text", "more"],
        }
    )
    assert guess_text_column(frame) == "feedback"


def test_text_column_only_looks_at_the_first_200_rows():
    frame = pd.DataFrame({"a": ["xx"] * 200 + ["x" * 1000] * 50, "b": ["xxx"] * 250})
    assert guess_text_column(frame) == "b"


def test_text_column_on_empty_frames():
    assert guess_text_column(pd.DataFrame()) is None
    assert guess_text_column(pd.DataFrame({"a": pd.Series([], dtype=str)})) is None


def test_text_column_handles_non_string_columns():
    frame = pd.DataFrame({"n": [123456789, 2], "t": ["abc", "de"]})
    assert guess_text_column(frame) == "n"


def test_id_column_requires_unique_complete_values():
    assert guess_id_column(pd.DataFrame({"id": ["1", "1"], "text": ["a", "b"]})) is None
    assert guess_id_column(pd.DataFrame({"ID": ["1", None], "text": ["a", "b"]})) is None
    assert guess_id_column(pd.DataFrame({"Key": ["001", "01"]})) == "Key"


def test_id_column_priority_and_fallback():
    frame = pd.DataFrame({"ref": ["a", "b"], "id": ["1", "2"], "index": ["x", "x"]})
    assert guess_id_column(frame) == "id"  # "index" first in priority, but not unique
    assert guess_id_column(pd.DataFrame({"text": ["a", "b"]})) is None
    assert guess_id_column(pd.DataFrame({"id": pd.Series([], dtype=str)})) is None


@pytest.mark.parametrize("name", ["case_id", "Ticket ID", "customer-id", "CaseID"])
def test_id_column_suffix_names(name):
    assert guess_id_column(pd.DataFrame({"text": ["a", "b"], name: ["1", "2"]})) == name


@pytest.mark.parametrize("name", ["VALID", "paid", "idea", "identity"])
def test_id_column_ignores_lookalikes(name):
    assert guess_id_column(pd.DataFrame({name: ["1", "2"]})) is None


def test_exact_id_name_beats_suffix_name():
    frame = pd.DataFrame({"case_id": ["1", "2"], "uid": ["a", "b"]})
    assert guess_id_column(frame) == "uid"


def test_gold_columns_name_prefixes_come_first():
    n = 10
    frame = pd.DataFrame(
        {
            "text": [f"text {i}" for i in range(n)],
            "tier": ["1", "2"] * 5,
            "Label_free": [f"free {i}" for i in range(n)],  # high cardinality, still named
            "y_true": ["a", "b"] * 5,
            "expected": ["x"] * n,
            "unrelated": [str(i) for i in range(n)],
        }
    )
    assert guess_gold_columns(frame, exclude=("text",)) == [
        "Label_free",
        "y_true",
        "expected",
        "tier",
    ]
    assert "y_true" not in guess_gold_columns(frame, exclude=("y_true",))


def test_gold_low_cardinality_needs_fewer_values_than_half_the_rows():
    frame = pd.DataFrame({"a": ["1", "2", "3", "1"], "b": ["x", "x", "y", "y"]})
    assert guess_gold_columns(frame) == []  # a: 3 >= 2, b: 2 >= 2
    frame = pd.DataFrame({"b": ["x", "x", "y", "y", "x"]})
    assert guess_gold_columns(frame) == ["b"]


def test_gold_limit_of_twenty_classes():
    rows = 100
    frame = pd.DataFrame(
        {
            "twenty": [str(i % 20) for i in range(rows)],
            "twenty_one": [str(i % 21) for i in range(rows)],
        }
    )
    assert guess_gold_columns(frame) == ["twenty"]


def test_gold_never_offers_an_empty_column():
    frame = pd.DataFrame({"gold": [None, None, None], "label": ["a", "a", "b"]})
    assert guess_gold_columns(frame) == ["label"]


def test_gold_and_group_on_empty_frames():
    assert guess_gold_columns(pd.DataFrame()) == []
    assert guess_group_columns(pd.DataFrame()) == []


def test_group_columns_need_two_to_twenty_values_with_repeats():
    rows = 40
    frame = pd.DataFrame(
        {
            "constant": ["same"] * rows,
            "pair": ["a", "b"] * 20,
            "twenty": [str(i % 20) for i in range(rows)],
            "twenty_one": [str(i % 21) for i in range(rows)],
            "with_gaps": ["x", None, "y", None] * 10,
        }
    )
    assert guess_group_columns(frame) == ["pair", "twenty", "with_gaps"]
    assert guess_group_columns(frame, exclude=("pair",)) == ["twenty", "with_gaps"]


def test_group_columns_skip_all_unique_columns():
    frame = pd.DataFrame({"text": ["a", "b", "c"], "tier": ["1", "1", "2"]})
    assert guess_group_columns(frame) == ["tier"]


def test_text_column_name_needs_a_filled_column_of_real_text():
    feedback = [
        "The parcel arrived late and the box was crushed",
        "Lovely staff and quick service overall",
        "It was fine I suppose, nothing special",
        "Never again, the worst support I have had",
    ]
    sparse = pd.DataFrame(
        {"Ticket ID": ["1", "2", "3", "4"], "Customer feedback": feedback,
         "Notes": [None, "call back", None, None]}
    )  # fmt: skip
    assert guess_text_column(sparse) == "Customer feedback"
    empty = sparse.assign(Notes=[None] * 4)
    assert guess_text_column(empty) == "Customer feedback"
    german = pd.DataFrame({"ID": ["1", "2", "3", "4"], "Kommentar": feedback,
                           "Note": ["1,5", "2", "3", "1"]})  # fmt: skip
    assert guess_text_column(german) == "Kommentar"
    # A well-filled notes column of real text still wins by name.
    notes = sparse.assign(Notes=["called, sorted", "agent apologised", "none", "refund sent"])
    assert guess_text_column(notes) == "Notes"


def test_guessers_skip_the_result_columns_of_an_exported_file():
    n = 4  # copies of five rows
    frame = pd.DataFrame(
        {
            "Ticket": [str(i) for i in range(5 * n)],
            "Feedback": [
                "parcel came late",
                "really nice staff",
                "ok then I guess",
                "bad bad service",
                "fine overall",
            ]
            * n,
            "Sentiment": ["negative", "positive", "neutral", "negative", "neutral"] * n,
            "OUTPUT": [f"A long generated summary of feedback number {i}" for i in range(5 * n)],
            "LABEL": ["negative", "positive", "neutral", "negative", "positive"] * n,
            "STATUS": ["generated", "generated", "cached", "error", "generated"] * n,
            "ERROR": ["", "", "", "CUDA out of memory, a long error message", ""] * n,
            "CORRECT": ["True", "True", "True", "", "False"] * n,
        }
    )
    # Read back the way a user would reload it.
    reloaded = read_table(results.to_csv_bytes(frame), name="export.csv").frame
    assert guess_text_column(reloaded) == "Feedback"
    assert guess_gold_columns(reloaded, exclude=("Feedback",)) == ["Sentiment", "LABEL"]
    # A second run on that export adds OUTPUT_2, LABEL_2, STATUS_2.
    again = reloaded.assign(
        OUTPUT_2=["x"] * 5 * n, LABEL_2=["positive", "negative"] * 10, STATUS_2=["cached"] * 5 * n
    )
    assert guess_gold_columns(again, exclude=("Feedback",)) == ["Sentiment", "LABEL", "LABEL_2"]


def test_own_label_and_status_columns_are_not_result_columns():
    frame = pd.DataFrame(
        {
            "text": [f"text number {i}" for i in range(8)],
            "OUTPUT": ["a", "b", "c", "d"] * 2,
            "LABEL": ["pos", "neg"] * 4,
            "STATUS": ["open", "closed", "open", "open"] * 2,  # not row statuses
        }
    )
    assert guess_gold_columns(frame, exclude=("text",)) == ["LABEL", "STATUS"]


def test_guessers_survive_unhashable_cells():
    frame = pd.DataFrame({"id": [[1], [2]], "tags": [["a"], ["b"]], "flag": ["x", "x"]})
    assert guess_id_column(frame) is None
    assert guess_gold_columns(frame) == []
    assert guess_group_columns(frame) == []
    assert guess_text_column(frame) in {"id", "tags"}


# --- folder listing ----------------------------------------------------------------------


def test_list_data_files_newest_first(tmp_path):
    names = ["old.csv", "middle.XLSX", "new.parquet", "same_b.jsonl", "same_a.json"]
    times = [1_000_000, 2_000_000, 3_000_000, 2_500_000, 2_500_000]
    for name, stamp in zip(names, times, strict=True):
        path = write(tmp_path, name, b"x")
        os.utime(path, (stamp, stamp))
    for ignored in ["notes.docx", "old.xls", ".hidden.csv", "~$lock.xlsx", "noext"]:
        write(tmp_path, ignored, b"x")
    (tmp_path / "folder.csv").mkdir()
    nested = tmp_path / "sub"
    nested.mkdir()
    write(nested, "deep.csv", b"x")

    found = list_data_files(tmp_path)
    assert [path.name for path in found] == [
        "new.parquet",
        "same_a.json",
        "same_b.jsonl",
        "middle.XLSX",
        "old.csv",
    ]
    assert all(path.parent == tmp_path for path in found)


def test_list_data_files_accepts_str(tmp_path):
    write(tmp_path, "a.txt", b"x")
    assert list_data_files(str(tmp_path)) == [tmp_path / "a.txt"]


def test_list_data_files_missing_or_not_a_folder(tmp_path):
    assert list_data_files(tmp_path / "missing") == []
    assert list_data_files(write(tmp_path, "file.csv", b"x")) == []
    assert list_data_files(tmp_path / "empty_dir_that_is_missing" / "deeper") == []


def test_list_data_files_empty_folder(tmp_path):
    assert list_data_files(tmp_path) == []


# --- fingerprint -------------------------------------------------------------------------


def test_fingerprint_is_short_hex_and_stable(sample):
    first = frame_fingerprint(sample)
    assert re.fullmatch(r"[0-9a-f]{16}", first)
    assert frame_fingerprint(sample) == first
    assert frame_fingerprint(sample.copy()) == first
    reread = read_table(SAMPLE.encode(), name="again.csv").frame
    assert frame_fingerprint(reread) == first


def test_fingerprint_changes_with_content(sample):
    base = frame_fingerprint(sample)
    edited = sample.copy()
    edited.loc[3, "TEXT"] = "changed"
    renamed = sample.rename(columns={"TEXT": "BODY"})
    reordered = sample.iloc[::-1]
    reindexed = sample.set_axis(range(100, 100 + len(sample)))
    gap = sample.copy()
    gap.loc[0, "RATIONALE"] = None
    fingerprints = {
        frame_fingerprint(frame) for frame in (edited, renamed, reordered, reindexed, gap)
    }
    assert base not in fingerprints
    assert len(fingerprints) == 5


def test_fingerprint_of_odd_frames():
    assert re.fullmatch(r"[0-9a-f]{16}", frame_fingerprint(pd.DataFrame()))
    assert frame_fingerprint(pd.DataFrame()) != frame_fingerprint(pd.DataFrame(columns=["a"]))
    listy = pd.DataFrame({"tags": [["a"], ["b"]]})
    assert frame_fingerprint(listy) == frame_fingerprint(listy.copy())
