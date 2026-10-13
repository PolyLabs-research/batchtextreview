"""Tests for result tables, exports, history records and the stats sentence."""

from __future__ import annotations

import io
import json
from datetime import datetime, timezone

import openpyxl
import pandas as pd
import pytest

import batchtextreview
from batchtextreview import results
from batchtextreview.engine import RunConfig, RunResult, RunStats
from batchtextreview.evaluation import evaluate
from batchtextreview.labels import INVALID
from batchtextreview.results import (
    build_results_frame,
    export_bytes,
    free_column_name,
    history_record,
    scorable_mask,
    summarise_stats,
    to_csv_bytes,
    to_excel_bytes,
)

CONFIG = RunConfig(
    model_id="org/model",
    model_label="Model",
    precision="fp16",
    template="Classify: {text}",
    max_new_tokens=8,
)
BOM = b"\xef\xbb\xbf"


def make_result(outputs, statuses=None, errors=None, **kwargs) -> RunResult:
    statuses = statuses or ["generated"] * len(outputs)
    errors = errors or [None] * len(outputs)
    stats = RunStats(
        total=len(outputs),
        blank=statuses.count("blank"),
        cached=statuses.count("cached"),
        generated=statuses.count("generated"),
        errors=statuses.count("error"),
        pending=statuses.count("pending"),
    )
    return RunResult(
        run_id=CONFIG.run_id,
        config=CONFIG,
        outputs=list(outputs),
        statuses=list(statuses),
        errors=list(errors),
        stats=stats,
        **kwargs,
    )


def stats_result(cancelled=False, aborted_reason=None, **counts) -> RunResult:
    stats = RunStats(**counts)
    n = stats.total
    return RunResult(
        run_id=CONFIG.run_id,
        config=CONFIG,
        outputs=[""] * n,
        statuses=["generated"] * n,
        errors=[None] * n,
        stats=stats,
        cancelled=cancelled,
        aborted_reason=aborted_reason,
    )


# --------------------------------------------------------------------------
# Column naming
# --------------------------------------------------------------------------


def test_free_column_name():
    frame = pd.DataFrame(columns=["text", "OUTPUT", "OUTPUT_2", 0])
    assert free_column_name(frame, "LABEL") == "LABEL"
    assert free_column_name(frame, "OUTPUT") == "OUTPUT_3"
    assert free_column_name(frame, "text") == "text_2"
    assert free_column_name(frame, "0") == "0_2"


def test_scorable_mask():
    result = make_result(
        ["a", "b", "", "", ""], ["generated", "cached", "blank", "error", "pending"]
    )
    assert scorable_mask(result) == [True, True, False, False, False]


# --------------------------------------------------------------------------
# build_results_frame
# --------------------------------------------------------------------------


def test_adds_output_and_status_keeping_original_columns():
    frame = pd.DataFrame({"id": [7, 8], "text": ["x", "y"]}, index=[10, 20])
    table = build_results_frame(frame, make_result(["X", "Y"], ["generated", "cached"]))
    assert list(table.columns) == ["id", "text", "OUTPUT", "STATUS"]
    assert list(table.index) == [0, 1]
    assert table["id"].tolist() == [7, 8]
    assert table["OUTPUT"].tolist() == ["X", "Y"]
    assert table["STATUS"].tolist() == ["generated", "cached"]
    # The caller's frame is untouched.
    assert list(frame.columns) == ["id", "text"]
    assert list(frame.index) == [10, 20]


def test_row_count_mismatch_raises():
    frame = pd.DataFrame({"text": ["a", "b"]})
    with pytest.raises(ValueError, match="rows"):
        build_results_frame(frame, make_result(["only one"]))


def test_column_collisions_get_free_names():
    frame = pd.DataFrame(
        {
            "text": ["a"],
            "OUTPUT": ["old output"],
            "OUTPUT_2": ["older"],
            "STATUS": ["old status"],
            "LABEL": ["old label"],
            "ERROR": ["old error"],
        }
    )
    result = make_result([""], ["error"], ["RuntimeError: boom"])
    table = build_results_frame(frame, result, labels=["yes", "no"])
    assert table["OUTPUT"].tolist() == ["old output"]
    assert table["OUTPUT_2"].tolist() == ["older"]
    assert table["STATUS"].tolist() == ["old status"]
    assert table["OUTPUT_3"].tolist() == [""]
    assert table["LABEL_2"].tolist() == [""]
    assert table["STATUS_2"].tolist() == ["error"]
    assert table["ERROR_2"].tolist() == ["RuntimeError: boom"]


def test_duplicate_input_columns_survive():
    frame = pd.DataFrame([["a", "b"]], columns=["text", "text"])
    table = build_results_frame(frame, make_result(["out"]))
    assert list(table.columns) == ["text", "text", "OUTPUT", "STATUS"]


def test_label_column_resolves_and_blanks_non_scorable_rows():
    outputs = ["Positive.", "maybe", "negative", "", "", "", ""]
    statuses = ["generated", "generated", "cached", "blank", "error", "pending", "generated"]
    errors = [None, None, None, None, "RuntimeError: x", None, None]
    frame = pd.DataFrame({"text": list("abcdefg")})
    table = build_results_frame(
        frame, make_result(outputs, statuses, errors), labels=["positive", "negative"]
    )
    assert table["LABEL"].tolist() == [
        "positive",
        INVALID,
        "negative",
        "",
        "",
        "",
        INVALID,  # the model produced nothing: scored as invalid too
    ]


@pytest.mark.parametrize("labels", [None, []])
def test_no_label_column_without_labels(labels):
    table = build_results_frame(pd.DataFrame({"t": ["a"]}), make_result(["x"]), labels=labels)
    assert "LABEL" not in table.columns


def test_error_column_only_when_needed():
    frame = pd.DataFrame({"text": ["a", "b"]})
    clean = build_results_frame(frame, make_result(["x", ""], ["generated", "pending"]))
    assert "ERROR" not in clean.columns

    failed = build_results_frame(
        frame, make_result(["x", ""], ["generated", "error"], [None, "ValueError: bad"])
    )
    assert failed["ERROR"].tolist() == ["", "ValueError: bad"]
    assert failed["OUTPUT"].tolist() == ["x", ""]


def test_no_note_column_without_truncation():
    frame = pd.DataFrame({"text": ["a", "b"]})
    result = make_result(["x", "y"], input_truncated=[False, False], hit_token_limit=[False] * 2)
    assert "NOTE" not in build_results_frame(frame, result).columns
    # Results made without the per-row flags.
    assert "NOTE" not in build_results_frame(frame, make_result(["x", "y"])).columns


def test_note_column_for_truncated_inputs_and_outputs():
    frame = pd.DataFrame({"text": list("abcd"), "NOTE": ["mine"] * 4})
    result = make_result(
        ["w", "x", "y", "z"],
        input_truncated=[True, False, True, False],
        hit_token_limit=[False, True, True, False],
    )
    table = build_results_frame(frame, result, labels=["x"])
    assert table["NOTE"].tolist() == ["mine"] * 4  # the input's own column is kept
    assert table["NOTE_2"].tolist() == [
        "input truncated",
        "output may be cut off",
        "input truncated; output may be cut off",
        "",
    ]
    assert list(table.columns) == ["text", "NOTE", "OUTPUT", "LABEL", "STATUS", "NOTE_2"]


def test_note_column_follows_the_error_column():
    frame = pd.DataFrame({"text": ["a", "b"]})
    result = make_result(
        ["", "x"], ["error", "generated"], ["E: x", None], input_truncated=[True, False]
    )
    table = build_results_frame(frame, result)
    assert list(table.columns) == ["text", "OUTPUT", "STATUS", "ERROR", "NOTE"]
    assert table["NOTE"].tolist() == ["input truncated", ""]


def test_note_flags_of_the_wrong_length_are_ignored():
    frame = pd.DataFrame({"text": ["a", "b"]})
    result = make_result(["x", "y"], input_truncated=[True], hit_token_limit=[True, True, True])
    assert "NOTE" not in build_results_frame(frame, result).columns


def test_column_order():
    frame = pd.DataFrame({"text": ["a"]})
    table = build_results_frame(
        frame, make_result([""], ["error"], ["E: x"]), labels=["a"], expand_json=True
    )
    assert list(table.columns) == [
        "text",
        "OUTPUT",
        "LABEL",
        "STATUS",
        "ERROR",
        "OUTPUT.valid_json",
    ]


def test_json_expansion():
    outputs = [
        json.dumps(
            {"people": ["Ann", "Bob"], "orgs": [], "meta": {"a": 1}, "n": 3, "ok": True, "x": None}
        ),
        '```json\n{"people": ["Cy", null, ""], "extra": "é"}\n```',
        "not json at all",
        "",
        "[1, 2]",
        "",
    ]
    statuses = ["generated", "cached", "generated", "blank", "generated", "error"]
    frame = pd.DataFrame({"text": list("abcdef")})
    table = build_results_frame(
        frame, make_result(outputs, statuses, [None] * 5 + ["E"]), expand_json=True
    )
    json_cols = [c for c in table.columns if c.startswith("OUTPUT.")]
    assert json_cols == [
        "OUTPUT.people",
        "OUTPUT.orgs",
        "OUTPUT.meta",
        "OUTPUT.n",
        "OUTPUT.ok",
        "OUTPUT.x",
        "OUTPUT.extra",
        "OUTPUT.valid_json",
    ]
    assert table["OUTPUT.people"].tolist() == ["Ann; Bob", "Cy", "", "", "", ""]
    assert table["OUTPUT.orgs"].tolist()[0] == ""
    assert table["OUTPUT.meta"].tolist()[0] == '{"a":1}'
    assert table["OUTPUT.n"].tolist()[0] == "3"
    assert table["OUTPUT.ok"].tolist()[0] == "true"
    assert table["OUTPUT.x"].tolist()[0] == ""
    assert table["OUTPUT.extra"].tolist() == ["", "é", "", "", "", ""]
    assert table["OUTPUT.valid_json"].tolist() == [True, True, False, None, True, None]
    # OUTPUT itself is unchanged.
    assert table["OUTPUT"].tolist() == outputs


def test_cut_off_or_malformed_json_object_is_not_valid_json():
    # The nested "people" list once came back as the row's JSON, so the row
    # read as valid JSON with every key column empty.
    outputs = [
        '{"people": ["Ann Lee", "Bob Ray"], "organisations": ["BPS"], "locations": ["Leeds", "York',
        '{"people": ["John"], "organisations": [], "locations": [], "dates": [],}',
        '{"people": ["Cy"], "organisations": [], "locations": [], "dates": []}',
    ]
    table = build_results_frame(
        pd.DataFrame({"t": [1, 2, 3]}),
        make_result(outputs, hit_token_limit=[True, False, False]),
        expand_json=True,
    )
    assert table["OUTPUT.valid_json"].tolist() == [False, False, True]
    assert table["OUTPUT.people"].tolist() == ["", "", "Cy"]
    assert table["NOTE"].tolist()[0] == "output may be cut off"


def test_json_list_of_objects_joined_as_compact_json():
    out = json.dumps({"items": [{"a": 1}, ["x", "y"], 2]})
    table = build_results_frame(pd.DataFrame({"t": [1]}), make_result([out]), expand_json=True)
    assert table["OUTPUT.items"].tolist() == ['{"a":1}; ["x","y"]; 2']


def test_json_columns_follow_renamed_output_column():
    frame = pd.DataFrame({"text": ["a"], "OUTPUT": ["old"], "OUTPUT.people": ["old people"]})
    table = build_results_frame(frame, make_result(['{"people": ["Z"]}']), expand_json=True)
    assert table["OUTPUT.people"].tolist() == ["old people"]
    assert table["OUTPUT_2.people"].tolist() == ["Z"]
    assert table["OUTPUT_2.valid_json"].tolist() == [True]


def test_json_key_named_valid_json_does_not_clobber_the_flag():
    table = build_results_frame(
        pd.DataFrame({"t": [1]}), make_result(['{"valid_json": "nope"}']), expand_json=True
    )
    assert table["OUTPUT.valid_json"].tolist() == [True]
    assert table["OUTPUT.valid_json_2"].tolist() == ["nope"]


def test_json_columns_are_capped():
    big = json.dumps({f"k{i}": i for i in range(results.MAX_JSON_COLUMNS + 10)})
    small = json.dumps({"k59": 1})
    table = build_results_frame(
        pd.DataFrame({"t": [1, 2]}), make_result([big, small]), expand_json=True
    )
    key_cols = [c for c in table.columns if c.startswith("OUTPUT.") and c != "OUTPUT.valid_json"]
    assert len(key_cols) == results.MAX_JSON_COLUMNS
    assert "OUTPUT.k59" in key_cols  # most frequent keys win


def test_no_json_columns_by_default():
    table = build_results_frame(pd.DataFrame({"t": [1]}), make_result(['{"a": 1}']))
    assert not [c for c in table.columns if str(c).startswith("OUTPUT.")]


def test_empty_frame():
    frame = pd.DataFrame({"text": pd.Series([], dtype=object)})
    table = build_results_frame(frame, make_result([]), labels=["a"], expand_json=True)
    assert len(table) == 0
    assert {"OUTPUT", "LABEL", "STATUS", "OUTPUT.valid_json"} <= set(table.columns)


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


def test_csv_has_bom_and_round_trips():
    frame = pd.DataFrame(
        {
            "id": [1, 2, 3],
            "text": ["café ☕", 'quote " and, comma', "line one\nline two"],
            "score": [0.5, -5.0, 1e-3],
            "jp": ["日本語", "NA", "None"],
        }
    )
    payload = to_csv_bytes(frame)
    assert payload.startswith(BOM)
    assert not payload.startswith(BOM + BOM)
    back = pd.read_csv(io.BytesIO(payload), encoding="utf-8-sig", keep_default_na=False)
    assert list(back.columns) == ["id", "text", "score", "jp"]
    assert back["id"].tolist() == [1, 2, 3]
    assert back["text"].tolist() == frame["text"].tolist()
    assert back["score"].tolist() == frame["score"].tolist()
    assert back["jp"].tolist() == ["日本語", "NA", "None"]


def test_csv_has_no_index_and_missing_values_are_empty():
    frame = pd.DataFrame({"a": ["x", None], "b": [1.0, float("nan")]}, index=[5, 9])
    text = to_csv_bytes(frame).decode("utf-8-sig")
    assert text.splitlines() == ["a,b", "x,1.0", ","]


def test_csv_neutralises_formula_like_text():
    frame = pd.DataFrame(
        {
            "text": [
                "=1+1",
                "+cmd|' /C calc'!A0",
                "-x",
                "@SUM(A1)",
                "\tindent",
                "-5",
                "+3.5",
                "-1,000.25",
                "plain = text",
                "-",
                "-5+cmd|x",
            ],
            "number": [-5] * 11,
        }
    )
    back = pd.read_csv(
        io.BytesIO(to_csv_bytes(frame)), encoding="utf-8-sig", keep_default_na=False, dtype=str
    )
    assert back["text"].tolist() == [
        "'=1+1",
        "'+cmd|' /C calc'!A0",
        "'-x",
        "'@SUM(A1)",
        "'\tindent",
        "-5",
        "+3.5",
        "-1,000.25",
        "plain = text",
        "-",
        "'-5+cmd|x",
    ]
    assert back["number"].tolist() == ["-5"] * 11
    assert frame["text"].iloc[0] == "=1+1"  # input untouched


def test_csv_neutralises_formula_like_headers():
    frame = pd.DataFrame([["a", "b", "c", 1]], columns=['=HYPERLINK("http://x")', "-x", "-5", 7])
    header = to_csv_bytes(frame).decode("utf-8-sig").splitlines()[0]
    assert header == '"\'=HYPERLINK(""http://x"")",\'-x,-5,7'
    assert frame.columns[0] == '=HYPERLINK("http://x")'  # input untouched


def test_csv_with_duplicate_column_names():
    frame = pd.DataFrame([["=x", "ok", 1]], columns=["t", "t", "n"])
    text = to_csv_bytes(frame).decode("utf-8-sig")
    assert text.splitlines() == ["t,t,n", "'=x,ok,1"]


def test_csv_with_string_and_category_dtypes():
    frame = pd.DataFrame(
        {
            "s": pd.Series(["=a", None], dtype="string"),
            "c": pd.Series(["=b", "ok"], dtype="category"),
        }
    )
    back = pd.read_csv(io.BytesIO(to_csv_bytes(frame)), encoding="utf-8-sig", dtype=str)
    assert back["s"].tolist()[0] == "'=a"
    assert pd.isna(back["s"].tolist()[1])
    assert back["c"].tolist() == ["'=b", "ok"]


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------


def load_sheet(payload: bytes):
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    return workbook, workbook.worksheets[0]


def test_xlsx_opens_with_frozen_bold_header():
    frame = pd.DataFrame({"id": [1, 2], "text": ["a", "b"]}, index=[3, 4])
    _, sheet = load_sheet(to_excel_bytes(frame))
    assert sheet.title == "results"
    assert sheet.freeze_panes == "A2"
    rows = list(sheet.iter_rows(values_only=True))
    assert rows == [("id", "text"), (1, "a"), (2, "b")]  # no index column
    assert sheet["A1"].font.b
    assert sheet.auto_filter.ref == "A1:B3"


def test_xlsx_never_writes_formulas():
    frame = pd.DataFrame({"=HEADER": ["=1+1", '=HYPERLINK("http://x")', "ok"]})
    _, sheet = load_sheet(to_excel_bytes(frame))
    assert sheet["A1"].value == "=HEADER"
    assert sheet["A1"].data_type == "s"
    assert sheet["A2"].value == "=1+1"
    assert sheet["A2"].data_type == "s"
    assert sheet["A3"].data_type == "s"


def test_xlsx_writes_error_code_text_as_text_and_reads_back():
    # "#N/A" from a failed VLOOKUP in the input used to become an Excel
    # error cell, which reading the export back turned into a missing value.
    from batchtextreview.data import read_table

    codes = ["#N/A", "#DIV/0!", "#VALUE!", "#NAME?", "#REF!", "#NULL!", "#NUM!"]
    frame = pd.DataFrame({"TEXT": codes, "OUTPUT": list(reversed(codes))})
    payload = to_excel_bytes(frame)
    _, sheet = load_sheet(payload)
    for row in sheet.iter_rows(min_row=2):
        assert [cell.data_type for cell in row] == ["s", "s"]
    reloaded = read_table(payload, name="results.xlsx").frame
    assert reloaded["TEXT"].tolist() == codes
    assert reloaded["OUTPUT"].tolist() == list(reversed(codes))


def test_xlsx_values_and_types():
    frame = pd.DataFrame(
        {
            "int": [1, 2],
            "float": [1.5, float("nan")],
            "bool": [True, False],
            "text": ["bell\x07 char", None],
            "when": [datetime(2024, 5, 1, 12, 0, tzinfo=timezone.utc), pd.NaT],
            "obj": [{"a": [1, 2]}, ["x"]],
            "big": [float("inf"), 2.0],
        }
    )
    _, sheet = load_sheet(to_excel_bytes(frame))
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == 1 and rows[0][1] == 1.5 and rows[0][2] is True
    assert rows[0][3] == "bell char"  # control character dropped
    assert rows[0][4] == datetime(2024, 5, 1, 12, 0)
    assert json.loads(rows[0][5]) == {"a": [1, 2]}
    assert rows[0][6] == "inf"
    assert rows[1][1] is None and rows[1][3] is None and rows[1][4] is None
    assert rows[1][2] is False


def test_xlsx_numbers_keep_their_precision():
    from decimal import Decimal

    import numpy as np

    frame = pd.DataFrame(
        {
            "id": pd.Series([1234567890123456789, 7], dtype="int64"),
            "flag": pd.Series([np.True_, np.False_], dtype=object),
            "dec": [Decimal("1.25"), Decimal("NaN")],
            "nullable": pd.Series([1, None], dtype="Int64"),
        }
    )
    _, sheet = load_sheet(to_excel_bytes(frame))
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "1234567890123456789"  # text: Excel would round it
    assert rows[1][0] == 7
    assert rows[0][1] is True and rows[1][1] is False
    assert rows[0][2] == 1.25
    assert rows[1][2] == "NaN"
    assert rows[0][3] == 1 and rows[1][3] is None


def test_xlsx_column_widths_capped():
    frame = pd.DataFrame({"short": ["a"], "long": ["x" * 500], "multi": ["abc\n" + "y" * 300]})
    _, sheet = load_sheet(to_excel_bytes(frame))
    widths = {c: sheet.column_dimensions[c].width for c in "ABC"}
    assert widths["A"] < 20
    assert widths["B"] == 80
    assert widths["C"] < 20  # only the first line counts
    assert all(w <= 80 for w in widths.values())


def test_xlsx_truncates_huge_cells():
    frame = pd.DataFrame({"text": ["z" * 40000]})
    _, sheet = load_sheet(to_excel_bytes(frame))
    assert len(sheet["A2"].value) == results.EXCEL_MAX_CELL_CHARS


def test_xlsx_sheet_name_sanitised():
    _, sheet = load_sheet(to_excel_bytes(pd.DataFrame({"a": [1]}), "a/b:c" + "x" * 40))
    assert len(sheet.title) <= 31
    assert "/" not in sheet.title and ":" not in sheet.title


def test_xlsx_empty_frame():
    _, sheet = load_sheet(to_excel_bytes(pd.DataFrame({"a": pd.Series([], dtype=object)})))
    assert list(sheet.iter_rows(values_only=True)) == [("a",)]


def test_export_bytes():
    frame = pd.DataFrame({"a": [1]})
    payload, mime = export_bytes(frame, "csv")
    assert mime == "text/csv" and payload.startswith(BOM)
    payload, mime = export_bytes(frame, "xlsx")
    assert mime == results.XLSX_MIME and payload[:2] == b"PK"
    with pytest.raises(ValueError):
        export_bytes(frame, "pdf")  # type: ignore[arg-type]


# --------------------------------------------------------------------------
# History record
# --------------------------------------------------------------------------


def labelled_report():
    preds = ["positive", "negative", "maybe", "", "positive", "positive"]
    golds = ["positive", "positive", "negative", "negative", None, "positive"]
    include = [True, True, True, False, True, True]
    return evaluate(preds, golds, allowed=["positive", "negative"], include=include)


def test_history_record_fields():
    report = labelled_report()
    result = make_result(["x"] * 6)
    record = history_record(
        report,
        result,
        preset_key="Sentiment",
        var="  ",
        dataset="reviews.csv",
        dataset_fingerprint="abc123",
        text_column="text",
        gold_column="gold",
        n_rows=6,
    )
    assert record == {
        "run_id": CONFIG.run_id,
        "model_id": "org/model",
        "model_label": "Model",
        "precision": "fp16",
        "preset": "Sentiment",
        "var": None,
        "labels": ["positive", "negative"],
        "dataset": "reviews.csv",
        "dataset_fingerprint": "abc123",
        "text_column": "text",
        "gold_column": "gold",
        "n_rows": 6,
        "n_scored": report.n_scored,
        "n_correct": report.n_correct,
        "accuracy": round(report.accuracy, 4),
        "macro_f1": round(report.macro_f1, 4),
        "kappa": None if report.kappa is None else round(report.kappa, 4),
        "n_invalid": report.n_pred_invalid,
        "n_excluded": 1,
        "n_gold_unscorable": 1,
        "source": "app",
        "selection": None,
        "rows_per_s": None,  # the result has no generation time
        "elapsed_s": 0.0,
        "app_version": batchtextreview.__version__,
        "template_edited": None,
        "preset_title": None,
    }
    assert report.n_scored == 4
    assert record["n_invalid"] == 1
    assert record["accuracy"] == 0.5
    for key in ("accuracy", "macro_f1", "kappa"):
        value = record[key]
        assert value is None or value == round(value, 4)


def history_kwargs(**overrides):
    kwargs = {
        "preset_key": "Sentiment",
        "var": None,
        "dataset": "d.csv",
        "dataset_fingerprint": "fp",
        "text_column": "text",
        "gold_column": "gold",
        "n_rows": 6,
    }
    kwargs.update(overrides)
    return kwargs


def test_history_record_provenance():
    result = make_result(["x"] * 6)
    result.stats.generation_s = 2.0
    result.stats.elapsed_s = 3.123456
    record = history_record(
        labelled_report(),
        result,
        **history_kwargs(
            selection=" First 6 of 10 rows ", template_edited=1, preset_title="Sentiment"
        ),
    )
    assert record["selection"] == "First 6 of 10 rows"
    assert record["template_edited"] is True
    assert record["preset_title"] == "Sentiment"
    assert record["rows_per_s"] == 3.0  # 6 generated rows in 2 s of generation
    assert record["elapsed_s"] == 3.1235
    assert record["app_version"] == batchtextreview.__version__


def test_history_record_provenance_blanks():
    result = make_result(["x"] * 6, ["cached"] * 6)
    result.stats.generation_s = 0.5  # nothing generated: no rate
    record = history_record(
        labelled_report(),
        result,
        **history_kwargs(selection="  ", template_edited=False, preset_title=""),
    )
    assert record["selection"] is None
    assert record["template_edited"] is False
    assert record["preset_title"] is None
    assert record["rows_per_s"] is None
    assert record["elapsed_s"] is None  # only the time it took to read the cache


def test_history_record_partly_cached_run_has_figures():
    result = make_result(["x"] * 6, ["cached"] * 5 + ["generated"])
    result.stats.generation_s = 0.5
    result.stats.elapsed_s = 0.75
    record = history_record(labelled_report(), result, **history_kwargs())
    assert record["rows_per_s"] == 2.0
    assert record["elapsed_s"] == 0.75


def test_rescoring_a_cached_run_keeps_its_history_figures(tmp_path):
    """The same run again, served from the cache, by a caller that passes no
    selection or preset title: the history line keeps the first run's
    throughput and provenance (and still gets the new metrics)."""
    from batchtextreview.store import Store

    store = Store(tmp_path / "btr.db")
    first = make_result(["x"] * 6)
    first.stats.generation_s = 0.1
    first.stats.elapsed_s = 0.0569
    store.log_evaluation(
        history_record(
            labelled_report(),
            first,
            **history_kwargs(
                selection="All 6 rows", template_edited=False, preset_title="Sentiment"
            ),
        )
    )
    cached = make_result(["x"] * 6, ["cached"] * 6)
    cached.stats.elapsed_s = 0.0009
    report = evaluate(["positive"] * 6, ["positive"] * 6, allowed=["positive", "negative"])
    store.log_evaluation(history_record(report, cached, **history_kwargs()))
    history = store.list_evaluations()
    assert len(history) == 1
    line = history.iloc[0]
    assert line["accuracy"] == 1.0
    assert line["rows_per_s"] == 60.0
    assert line["elapsed_s"] == 0.0569
    assert line["selection"] == "All 6 rows"
    assert line["template_edited"] is False
    assert line["preset_title"] == "Sentiment"


def test_history_record_exact_mode_has_no_labels():
    report = evaluate(["Paris", "Rome"], ["paris", "Oslo"], allowed=None)
    record = history_record(
        report,
        make_result(["Paris", "Rome"]),
        preset_key="Custom",
        var="capital of {x}",
        dataset=None,
        dataset_fingerprint=None,
        text_column="t",
        gold_column="g",
        n_rows=2,
        source="cli",
    )
    assert record["labels"] is None
    assert record["var"] == "capital of {x}"
    assert record["accuracy"] == 0.5
    assert record["macro_f1"] is None
    assert record["source"] == "cli"


def test_history_record_logs_into_real_store(tmp_path):
    from batchtextreview.store import Store

    try:
        store = Store(tmp_path / "btr.db")
    except NotImplementedError:
        pytest.skip("Store not implemented yet")
    record = history_record(
        labelled_report(),
        make_result(["x"] * 6),
        preset_key="Sentiment",
        var=None,
        dataset="d.csv",
        dataset_fingerprint="fp",
        text_column="text",
        gold_column="gold",
        n_rows=6,
    )
    first = store.log_evaluation(record)
    again = store.log_evaluation(record)  # same identity: updated, not duplicated
    assert first == again
    history = store.list_evaluations()
    assert len(history) == 1
    assert history["labels"].iloc[0] == ["positive", "negative"]
    assert history["accuracy"].iloc[0] == 0.5

    result = make_result(["x"] * 6)
    result.stats.generation_s = 3.0
    result.stats.elapsed_s = 4.5
    store.log_evaluation(
        history_record(
            labelled_report(),
            result,
            **history_kwargs(
                dataset_fingerprint="fp2",
                selection="All 6 rows",
                template_edited=True,
                preset_title="Sentiment",
            ),
        )
    )
    line = store.list_evaluations().iloc[0]
    assert line["selection"] == "All 6 rows"
    assert line["rows_per_s"] == 2.0
    assert line["elapsed_s"] == 4.5
    assert line["app_version"] == batchtextreview.__version__
    assert line["template_edited"] is True
    assert line["preset_title"] == "Sentiment"


# --------------------------------------------------------------------------
# Stats sentence
# --------------------------------------------------------------------------


def test_summary_example_sentence():
    result = stats_result(total=120, generated=80, cached=30, blank=10, duplicates=12)
    assert summarise_stats(result) == (
        "120 rows: 80 generated, 30 from cache, 10 blank. 12 duplicate rows were computed once."
    )


def test_summary_singular_and_omissions():
    assert summarise_stats(stats_result(total=1, generated=1)) == "1 row: 1 generated."
    text = summarise_stats(stats_result(total=3, generated=2, blank=1, duplicates=1))
    assert text == "3 rows: 2 generated, 1 blank. 1 duplicate row was computed once."
    assert "from cache" not in text


def test_summary_zero_rows():
    assert summarise_stats(stats_result(total=0)) == "0 rows."


def test_summary_errors_and_pending():
    result = stats_result(total=1500, generated=1200, errors=3, pending=297)
    assert summarise_stats(result) == "1,500 rows: 1,200 generated, 3 failed, 297 not run."


def test_summary_cancelled():
    text = summarise_stats(stats_result(total=10, generated=4, pending=6, cancelled=True))
    assert text == "10 rows: 4 generated, 6 not run. The run was cancelled."


def test_summary_aborted():
    text = summarise_stats(
        stats_result(
            total=5,
            errors=2,
            pending=3,
            aborted_reason="RuntimeError: CUDA error.\nDetails follow",
        )
    )
    assert text == (
        "5 rows: 2 failed, 3 not run. Stopped early: RuntimeError: CUDA error. Details follow."
    )


def test_summary_long_abort_reason_is_shortened():
    text = summarise_stats(stats_result(total=1, pending=1, aborted_reason="x" * 1000))
    assert len(text) < 260


def test_summary_mentions_truncated_inputs_and_token_limit():
    text = summarise_stats(
        stats_result(total=10, generated=10, truncated_inputs=3, hit_token_limit=5)
    )
    assert text == (
        "10 rows: 10 generated. 3 texts were longer than 8,000 characters and were cut "
        "before sending. 5 outputs reached the 8-token limit and may be cut off - raise "
        "Max output tokens."
    )


def test_summary_token_limit_advice_names_the_callers_setting():
    result = stats_result(total=2, generated=2, hit_token_limit=2)
    text = summarise_stats(result, token_setting="--max-new-tokens")
    assert text.endswith("may be cut off - raise --max-new-tokens.")
    assert "Max output tokens" not in text
    assert summarise_stats(result).endswith("raise Max output tokens.")


def test_summary_truncation_singular():
    text = summarise_stats(
        stats_result(total=1, generated=1, truncated_inputs=1, hit_token_limit=1)
    )
    assert "1 text was longer than 8,000 characters and was cut before sending." in text
    assert "1 output reached the 8-token limit and may be cut off" in text


def test_summary_truncation_comes_before_cancellation():
    text = summarise_stats(
        stats_result(total=4, generated=2, pending=2, truncated_inputs=2, cancelled=True)
    )
    assert text.endswith("were cut before sending. The run was cancelled.")
    assert "token limit" not in text


def _demo_run(texts, store, max_new_tokens=8, cancel=None):
    import threading

    from batchtextreview import engine
    from batchtextreview.backends.fake import FakeBackend
    from batchtextreview.backends.manager import ModelManager
    from batchtextreview.models import DEMO_MODEL

    config = RunConfig.build(DEMO_MODEL, "Summarise.\nTEXT:\n{text}", max_new_tokens)
    manager = ModelManager(lambda model_id, precision: FakeBackend(model_id, precision))
    event = threading.Event()
    if cancel:
        event.set()
    return engine.execute_run(config, texts, store=store, manager=manager, cancel=event)


def test_summary_claims_no_duplicates_for_runs_that_computed_nothing(tmp_path):
    """Regression: "98 duplicate rows were computed once" for a run
    cancelled before it started, and for a fully cached re-run."""
    from batchtextreview.store import Store

    store = Store(tmp_path / "btr.db")
    texts = ["great"] * 50 + ["awful"] * 50
    cancelled = summarise_stats(_demo_run(texts, store, cancel=True))
    assert cancelled == "100 rows: 100 not run. The run was cancelled."
    first = summarise_stats(_demo_run(texts, store))
    assert first == "100 rows: 100 generated. 98 duplicate rows were computed once."
    assert summarise_stats(_demo_run(texts, store)) == "100 rows: 100 from cache."


def test_cut_off_note_and_advice_survive_a_cached_re_run(tmp_path):
    """Regression: re-running (or resuming) from the saved outputs dropped
    the NOTE column and the "may be cut off" advice."""
    from batchtextreview.store import Store

    store = Store(tmp_path / "btr.db")
    texts = ["the parcel arrived two weeks late and the box was crushed", "fine"]
    frame = pd.DataFrame({"t": texts})
    for expected in ("generated", "cached"):
        result = _demo_run(texts, store, max_new_tokens=3)
        assert result.statuses == [expected, expected]
        assert result.hit_token_limit == [True, False]
        assert "1 output reached the 3-token limit and may be cut off" in summarise_stats(result)
        table = build_results_frame(frame, result)
        assert table["NOTE"].tolist() == ["output may be cut off", ""]
